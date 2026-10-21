"""Smoke tests for index.html, run in a real headless Chromium.

    pip install playwright && python -m playwright install --with-deps chromium
    python tests/smoke_test.py

GitHub is mocked (no network, no real token); Google Fonts is blocked so the
tests don't depend on the network either.
"""
import base64
import tempfile

import openpyxl
import re
import datetime
import json
import pathlib
import sys

from playwright.sync_api import sync_playwright

URL = (pathlib.Path(__file__).resolve().parent.parent / "index.html").as_uri()
TODAY = datetime.date.today()
TOMORROW = (TODAY + datetime.timedelta(days=1)).isoformat()
failures = []


def check(cond, msg):
    print(("ok   " if cond else "FAIL ") + msg)
    if not cond:
        failures.append(msg)


def task(tid, title, archived=False, updated=1000, **kw):
    t = {"id": tid, "title": title, "description": "", "client": "", "project": "", "urgency": "moyenne",
         "status": "a_faire", "due": "", "reminder": "", "link": "", "recur": None,
         "subtasks": [{"text": "a", "done": False}], "archived": archived, "archivedAt": None,
         "deleted": False, "createdAt": 1, "updatedAt": updated}
    t.update(kw)
    return t


def new_page(browser, seed=None, **ctx):
    page = browser.new_page(**ctx)
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.route("**/fonts.googleapis.com/**", lambda r: r.abort())
    page.route("**/fonts.gstatic.com/**", lambda r: r.abort())
    # older checks were written for the grouped contacts view; the default (list) is checked on its own
    page.add_init_script("try{ if(!sessionStorage.getItem('cvs')){ sessionStorage.setItem('cvs','1'); if(localStorage.getItem('contacts-view')===null) localStorage.setItem('contacts-view','groupe'); } }catch(e){}")
    if seed:
        page.add_init_script(
            "if(!sessionStorage.getItem('seeded')){ for (const [k,v] of Object.entries(" + json.dumps(seed) + ")) localStorage.setItem(k,v); sessionStorage.setItem('seeded','1'); }")
    return page, errors


# the app keeps its data in IndexedDB (in-memory copy: window.storage.peek); localStorage only for settings
PEEK = "(k=>(window.storage&&window.storage.peek)?window.storage.peek(k):localStorage.getItem('todo_'+k))"


def postit_edit(page):
    """The post-it shows its lines as a checklist; a click on a line's text edits it (cursor at the end of that line)."""
    if page.is_visible("#postit-list"): page.locator("#postit-list .pi-text").last.click()
    else: postit_edit(page)


def stored(page, key):
    page.wait_for_function("!!window.storage")
    page.evaluate("async () => { await window.storage.ready; }")
    return page.evaluate(f"JSON.parse({PEEK}('{key}')||'null')")


def mock_github(page, files, put_ok=True, large=()):
    """files: {'tasks.json': [...], ...} — server state, updated on successful PUTs.
    large: files GitHub reports as > 1 MB (metadata without content, raw fetch needed)."""
    state = {"puts": 0, "large": large}

    def handler(route, request):
        name = request.url.rsplit("/", 1)[-1]
        if name not in files and request.method == "GET":   # a PUT creates the file, like GitHub
            return route.fulfill(status=404, body="{}")
        if request.method == "GET":
            if "raw" in request.headers.get("accept", ""):
                return route.fulfill(status=200, content_type="application/vnd.github.raw", body=json.dumps(files[name]))
            if name in state.get("large", ()):  # > 1 MB: GitHub sends the metadata without content
                body = {"sha": "sha-" + name, "content": "", "encoding": "none"}
            else:
                body = {"sha": "sha-" + name, "encoding": "base64", "content": base64.b64encode(json.dumps(files[name]).encode()).decode()}
            return route.fulfill(status=200, content_type="application/json", body=json.dumps(body))
        state["puts"] += 1
        if not put_ok:
            return route.fulfill(status=500, body="{}")
        files[name] = json.loads(base64.b64decode(json.loads(request.post_data)["content"]).decode())
        route.fulfill(status=200, content_type="application/json", body=json.dumps({"content": {"sha": "s" + str(state["puts"])}}))

    page.route("**/api.github.com/repos/o/r/contents/data/*", handler)
    return state


SYNC_ON = {"todo_sync-config": json.dumps({"owner": "o", "repo": "r", "token": "t", "enabled": True})}

with sync_playwright() as p:
    browser = p.chromium.launch()

    # --- Tasks: create, validate (archive), delete -> Corbeille -> restore
    page, errors = new_page(browser)
    page.goto(URL)
    page.click('[data-tab="tasks"]')
    page.fill("#quick-add-input", "Tâche A"); page.press("#quick-add-input", "Enter")
    page.fill("#quick-add-input", "Tâche B"); page.press("#quick-add-input", "Enter")
    page.click("#quick-add-more"); page.click('#add-section .qd-chips button:text-is("Demain")')
    tomorrow = (TODAY + datetime.timedelta(days=1)).isoformat()
    check(page.input_value("#f-due") == tomorrow and page.locator("#add-section .qd-chips button.on").inner_text() == "Demain",
          "date field: one click on « Demain » fills the deadline")
    page.click("#reminder-toggle"); page.click('#reminder-row .qd-chips button:text-is("Demain 9 h")')
    check(page.input_value("#f-reminder") == tomorrow + "T09:00", "reminder: « Demain 9 h » fills date and time")
    page.click('#add-section .qd-chips .qd-clear >> nth=0')
    check(page.input_value("#f-due") == "" and page.evaluate("getComputedStyle(document.documentElement).colorScheme") == "dark",
          "date field: ✕ clears it; native pickers drawn for a dark page (white calendar icon)")
    page.evaluate("document.getElementById('f-reminder').value=''")
    page.click("#quick-add-more")
    page.locator("#task-list .task-line", has_text="Tâche A").locator(".task-check").click()
    check([t["archived"] for t in stored(page, "todo-widget-tasks") if t["title"] == "Tâche A"] == [True], "validated task is archived")
    check(page.locator("#task-list .task-line", has_text="Tâche A").count() == 0, "validated task leaves the open list")
    page.on("dialog", lambda d: d.accept())
    page.click('[data-tab="archive"]')
    row = page.locator("#archive-list .task-line", has_text="Tâche A")
    row.hover(); row.locator(".del-btn").click()
    page.wait_for_timeout(300)
    check(page.locator("#archive-list .task-line", has_text="Tâche A").count() == 0, "deleted task leaves the archive")
    page.click('#archive-filters [data-filter="corbeille"]')
    check(page.locator("#archive-list .task-line", has_text="Tâche A").count() == 1, "deleted task shows in Archive > Corbeille")
    page.locator("#archive-list .task-line", has_text="Tâche A").locator(".task-check").click()
    check([t["deleted"] for t in stored(page, "todo-widget-tasks") if t["title"] == "Tâche A"] == [False], "Corbeille restore un-deletes")
    check(page.locator("#archive-list .task-line", has_text="Tâche A").count() == 0, "restored task leaves the Corbeille")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Sync: a validation made right before closing the app (never pushed) survives the next pulls
    for put_ok in (True, False):
        files = {"tasks.json": [task("t1", "X", archived=False)]}
        seed = dict(SYNC_ON, **{"todo_todo-widget-tasks": json.dumps([task("t1", "X", archived=True, updated=2000)]),
                                "todo_sync-tasks-base": json.dumps([task("t1", "X", archived=False)])})
        page, errors = new_page(browser, seed)
        state = mock_github(page, files, put_ok=put_ok)
        page.goto(URL)
        page.wait_for_timeout(4500)  # pull + 3s debounced push
        page.click('[data-tab="lists"]'); page.click("#sync-now-btn"); page.wait_for_timeout(1500)
        check(stored(page, "todo-widget-tasks")[0]["archived"] is True, f"unpushed validation survives sync (put_ok={put_ok})")
        if put_ok:
            check(files["tasks.json"][0]["archived"] is True, "unpushed validation reaches the server")
        check(not errors, f"no page errors ({errors})")
        page.close()

    # --- widget.html uses the same three-way sync (an older copy elsewhere must not bring a task back)
    files = {"tasks.json": [task("t1", "X", archived=False, updated=3000, title2="ignored")]}
    files["tasks.json"][0]["description"] = "edited on the phone"
    seed = dict(SYNC_ON, **{"todo_todo-widget-tasks": json.dumps([task("t1", "X", archived=True, updated=2000)]),
                            "todo_sync-tasks-base": json.dumps([task("t1", "X", archived=False)])})
    page, errors = new_page(browser, seed)
    mock_github(page, files)
    page.goto(URL.replace("index.html", "widget.html")); page.wait_for_timeout(4500)
    t1 = stored(page, "todo-widget-tasks")[0]
    check(t1["archived"] is True and t1["description"] == "edited on the phone", "widget: validation kept and remote edit merged")
    check(files["tasks.json"][0]["archived"] is True, "widget: merged result pushed")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- widget mode: bubble, small task panel, same code as the app
    page, errors = new_page(browser, {"todo_todo-widget-tasks": json.dumps([task("w1", "Envoyer ARPE")])}, viewport={"width": 1200, "height": 760})
    page.goto(URL.replace("index.html", "widget.html")); page.wait_for_timeout(500)
    check("widget" in page.url and page.is_visible("#widget-bubble") and not page.is_visible("#todo-panel") and not page.is_visible(".app-sidebar"),
          "widget: widget.html opens the app's widget mode, closed on its bubble")
    check(page.inner_text("#widget-badge") == "1", "widget: open tasks counted on the bubble")
    page.click("#widget-bubble"); page.wait_for_timeout(150)
    check(page.is_visible("#todo-panel") and page.evaluate("document.activeElement.id") == "quick-add-input" and page.is_visible("#task-list .task-line:has-text('Envoyer ARPE')"),
          "widget: the bubble opens the task panel, ready to type")
    page.keyboard.type("Rappeler Paul"); page.keyboard.press("Enter")
    page.wait_for_function("document.getElementById('widget-badge').textContent==='2'", timeout=10000)
    check(page.inner_text("#widget-badge") == "2" and any(t["title"] == "Rappeler Paul" for t in stored(page, "todo-widget-tasks")), "widget: a task added from the widget is saved")
    page.keyboard.press("Escape"); page.wait_for_timeout(100)
    check(not page.is_visible("#todo-panel"), "widget: Escape folds the panel back into the bubble")
    b = page.locator("#widget-bubble").bounding_box()
    page.mouse.move(b["x"] + 20, b["y"] + 20); page.mouse.down(); page.mouse.move(300, 300, steps=8); page.mouse.up(); page.wait_for_timeout(100)
    check(not page.is_visible("#todo-panel") and abs(page.locator("#widget-bubble").bounding_box()["x"] - 280) < 6, "widget: the bubble can be dragged without opening the panel")
    page.reload(); page.wait_for_timeout(400)
    check(abs(page.locator("#widget-bubble").bounding_box()["x"] - 280) < 6, "widget: bubble position remembered")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Two tabs (app + widget): a change in one is picked up by the other, never overwritten
    ctx = browser.new_context()
    a = ctx.new_page(); b = ctx.new_page()
    for pg_ in (a, b):
        pg_.route("**/fonts.g*/**", lambda r: r.abort())
    a.goto(URL); a.click('[data-tab="tasks"]'); a.fill("#quick-add-input", "Onglet A"); a.press("#quick-add-input", "Enter")
    a.wait_for_function(f"({PEEK}('todo-widget-tasks')||'').includes('Onglet A')")
    b.goto(URL); b.click('[data-tab="tasks"]'); b.fill("#quick-add-input", "Onglet B"); b.press("#quick-add-input", "Enter")
    a.wait_for_function(f"({PEEK}('todo-widget-tasks')||'').includes('Onglet B')", timeout=10000)
    a.fill("#quick-add-input", "Encore A"); a.press("#quick-add-input", "Enter")
    try: b.wait_for_function(f"['Encore A','Onglet A','Onglet B'].every(x=>({PEEK}('todo-widget-tasks')||'').includes(x))", timeout=10000)
    except Exception: pass
    titles = sorted(t["title"] for t in b.evaluate(f"JSON.parse({PEEK}('todo-widget-tasks'))"))
    check(titles == ["Encore A", "Onglet A", "Onglet B"], f"two open tabs don't overwrite each other ({titles})")
    # the share page (older code writing straight into localStorage) while the app is open: merged, nothing lost
    b.goto(URL.replace("index.html", "share-target.html") + "?title=Partag%C3%A9+depuis+le+t%C3%A9l%C3%A9phone")
    a.wait_for_function(f"({PEEK}('todo-widget-tasks')||'').includes('Partagé') && localStorage.getItem('todo_todo-widget-tasks')===null", timeout=10000)
    titles = sorted(t["title"] for t in a.evaluate(f"JSON.parse({PEEK}('todo-widget-tasks'))"))
    check(titles == ["Encore A", "Onglet A", "Onglet B", "Partagé depuis le téléphone"] and a.evaluate("localStorage.getItem('todo_todo-widget-tasks')") is None,
          f"storage: a task written by the share page joins the open app's data, then leaves localStorage ({titles})")
    ctx.close()

    # --- Reminders: a reminder already shown doesn't fire again after reopening the app
    past = (datetime.datetime.now() - datetime.timedelta(minutes=5)).strftime("%Y-%m-%dT%H:%M")
    page, errors = new_page(browser, {"todo_todo-widget-tasks": json.dumps([task("r1", "Rappel test", reminder=past)])})
    page.add_init_script("window.__notifs=[]; window.Notification=function(t){ window.__notifs.push(t); }; Notification.permission='granted';")
    page.goto(URL); page.wait_for_timeout(500)
    first = page.evaluate("window.__notifs.length")
    page.reload(); page.wait_for_timeout(500)
    check(first == 1 and page.evaluate("window.__notifs.length") == 0, "reminder shown once, not again after reload")
    page.close()

    # --- Sync: a failed push must not make the next merge revert the change
    files = {"tasks.json": [task("t1", "X")]}
    seed = dict(SYNC_ON, **{"todo_todo-widget-tasks": json.dumps([task("t1", "X")]), "todo_sync-tasks-base": json.dumps([task("t1", "X")])})
    page, errors = new_page(browser, seed)
    mock_github(page, files, put_ok=False)
    page.goto(URL)
    page.click('[data-tab="tasks"]')
    page.locator("#task-list .task-check").first.click()
    page.wait_for_timeout(4000)
    page.click('[data-tab="lists"]'); page.click("#sync-now-btn"); page.wait_for_timeout(1500)
    check(stored(page, "todo-widget-tasks")[0]["archived"] is True, "failed push doesn't revert validation")
    page.close()

    # --- Post-it: type, persists, line -> task for tomorrow, minimise / reopen, legacy notes folded in
    legacy = [{"id": "n1", "text": "Rappeler Sophie", "kind": "demain", "forDate": TOMORROW, "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-notes": json.dumps(legacy)}, viewport={"width": 1300, "height": 850})
    page.goto(URL)
    check(page.is_visible("#postit"), "post-it open by default on desktop")
    check(page.input_value("#postit-text") == "Demain : Rappeler Sophie", "old multi-post-it notes folded into the post-it")
    postit_edit(page); page.keyboard.press("End")
    page.keyboard.type("\nEnvoyer offre PPA\nidée : webinar ACC")
    # the note is saved after a short pause in typing: wait for the save itself, not a fixed delay
    page.wait_for_function(f"({PEEK}('todo-widget-notes')||'').includes('webinar ACC')", timeout=10000)
    page.reload()
    check(page.input_value("#postit-text") == "Demain : Rappeler Sophie\nEnvoyer offre PPA\nidée : webinar ACC", "post-it text persists after reload")
    postit_edit(page)
    page.evaluate("const t=document.getElementById('postit-text'); const i=t.value.indexOf('Envoyer'); t.setSelectionRange(i+3,i+3);")
    page.click("#postit-to-task"); page.wait_for_timeout(300)
    tk = [t for t in stored(page, "todo-widget-tasks") if t["title"] == "Envoyer offre PPA"]
    check(len(tk) == 1 and tk[0]["due"] == TOMORROW and tk[0]["urgency"] == "haute", "line → task urgent for tomorrow")
    check("Envoyer offre PPA" not in page.input_value("#postit-text"), "converted line removed from post-it")
    page.click("#postit-min")
    check(not page.is_visible("#postit") and page.is_visible("#postit-tab"), "minimise shows the small tab")
    page.keyboard.press("p"); page.wait_for_timeout(100)
    check(page.is_visible("#postit") and page.evaluate("document.activeElement.id") == "postit-text", "'p' reopens the post-it with focus")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Storage: data moved from localStorage into IndexedDB, failures and space shown
    page, errors = new_page(browser, {"todo_todo-widget-tasks": json.dumps([task("m1", "Ancienne tâche")])}, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.click('[data-tab="tasks"]')
    check(page.is_visible("#task-list .task-line:has-text('Ancienne tâche')") and page.evaluate("localStorage.getItem('todo_todo-widget-tasks')") is None
          and page.evaluate("window.storage.mode") == "idb", "storage: existing data moved into the browser database, the old copy removed")
    page.reload(); page.click('[data-tab="tasks"]')
    check(page.is_visible("#task-list .task-line:has-text('Ancienne tâche')"), "storage: data still there after reload (read from the database)")
    page.click('[data-tab="lists"]'); page.wait_for_function("document.getElementById('storage-gauge').textContent.includes('Tâches')", timeout=5000)
    check("Base de données du navigateur" in page.inner_text("#storage-gauge"), "storage: gauge shows the space used, by kind of data")
    page.evaluate("() => { IDBObjectStore.prototype.put = function(){ throw new DOMException('full', 'QuotaExceededError'); }; }")
    page.click('[data-tab="tasks"]'); page.fill("#quick-add-input", "Tâche de trop"); page.press("#quick-add-input", "Enter")
    page.wait_for_selector("#storage-alert:not([hidden])", timeout=10000)
    check(page.is_visible("#storage-alert") and "plein" in page.inner_text("#storage-alert"), "storage: a failed save shows a red banner at once")
    page.click("#storage-alert-see")
    check(not page.is_visible("#storage-alert") and page.is_visible("#storage-section"), "storage: the banner leads to the space gauge")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Global search (Ctrl+K): everything in one list, Enter alone still creates the task
    gs_seed = {"todo_todo-widget-tasks": json.dumps([task("g1", "Envoyer ARPE Durand", client="Solaris"), task("g2", "Vieille relance Durand", archived=True)]),
               "todo_todo-widget-contacts": json.dumps([{"id": "gc1", "firstName": "Paul", "lastName": "Durand", "role": "Gérant", "company": "Solaris", "email": "p@solaris.fr",
                                                          "phone1": "06 12 00 00 99", "phone2": "", "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]),
               "todo_todo-widget-opportunities": json.dumps([{"id": "go1", "kind": "installation", "company": "Solaris", "centrale": "Toit Durand", "title": "Toit Durand",
                                                              "stage": "collecte", "kwc": "250", "createdAt": 1, "updatedAt": 1, "deleted": False}])}
    page, errors = new_page(browser, gs_seed, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.wait_for_timeout(300)
    page.keyboard.press("Control+k")
    check(page.is_visible("#gs-overlay") and page.evaluate("document.activeElement.id") == "gs-input", "search: Ctrl+K opens the global search")
    page.keyboard.type("durand")
    groups = page.eval_on_selector_all("#gs-results .gs-group", "els=>els.map(e=>e.textContent)")
    check(groups == ["Tâches", "Contacts", "Pipeline"] and page.locator("#gs-results .gs-item").count() == 5,
          f"search: tasks (archived too), contacts and pipeline found at once ({groups})")
    page.fill("#gs-input", "0612000099")
    check("Paul Durand" in page.inner_text("#gs-results"), "search: a contact found by phone number")
    page.locator("#gs-results .gs-item", has_text="Paul Durand").click(); page.wait_for_timeout(200)
    check(page.is_hidden("#gs-overlay") and page.is_visible("#tab-contacts") and page.is_visible("#contacts-list .contact-info-card, #contacts-list .contact-detail-links"),
          "search: choosing a contact opens its card in Contacts")
    page.keyboard.press("Control+k"); page.keyboard.type("toit"); page.keyboard.press("ArrowDown"); page.keyboard.press("Enter"); page.wait_for_timeout(300)
    check(page.is_visible("#opp-detail-overlay") and "Toit Durand" in page.inner_text("#opp-detail-title"), "search: arrows + Enter open the opportunity")
    page.keyboard.press("Escape"); page.wait_for_timeout(100)
    if page.is_visible("#opp-detail-overlay"): page.click("#opp-detail-actions >> text=Fermer")
    page.keyboard.press("Control+k"); page.keyboard.type("Appeler le notaire"); page.keyboard.press("Enter"); page.wait_for_timeout(200)
    check(any(t["title"] == "Appeler le notaire" for t in stored(page, "todo-widget-tasks")), "search: Enter without choosing still creates the task (old Ctrl+K habit)")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Company sheet: contacts, pipeline, tasks, what is coming up and the history of one client
    now_ms = int(datetime.datetime.now().timestamp() * 1000)
    late_day = (TODAY - datetime.timedelta(days=2)).isoformat()
    cs_seed = {"todo_todo-widget-tasks": json.dumps([task("s1", "Relancer pour le CR", client="Solaris", due=late_day),
                                                    task("s2", "Envoyer la proposition", client="Solaris", archived=True, archivedAt=now_ms - 86400000)]),
               "todo_todo-widget-contacts": json.dumps([{"id": "sc1", "firstName": "Paul", "lastName": "Durand", "role": "Gérant", "company": "Solaris", "email": "p@solaris.fr",
                                                        "phone1": "06 12 00 00 99", "phone2": "", "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]),
               "todo_todo-widget-meetings": json.dumps([{"id": "sm1", "contactId": "sc1", "label": "Salon EnerGaïa", "date": (TODAY - datetime.timedelta(days=5)).isoformat(),
                                                        "createdAt": 1, "updatedAt": 1, "deleted": False}]),
               "todo_todo-widget-opportunities": json.dumps([{"id": "so1", "kind": "installation", "company": "Solaris", "centrale": "Toit Durand", "title": "Toit Durand", "stage": "negociation",
                                                             "kwc": "1200", "createdAt": now_ms - 9 * 86400000, "updatedAt": 1, "deleted": False,
                                                             "stageLog": [{"stage": "collecte", "at": now_ms - 9 * 86400000}, {"stage": "negociation", "at": now_ms - 3 * 86400000}]}])}
    page, errors = new_page(browser, cs_seed, viewport={"width": 1300, "height": 850})
    page.goto(URL); page.click('[data-tab="contacts"]')
    page.locator(".company-group-header", has_text="Solaris").locator(".co-sheet-btn").click()
    sheet = page.inner_text("#co-sheet-overlay")
    check(page.is_visible("#co-sheet-overlay") and "1 en retard" in sheet and "1,2 MWc" in sheet and "Toit Durand" in sheet and "Paul Durand" in sheet,
          "company sheet: contacts, pipeline (MWc), open tasks with the late one, on one screen")
    hist = page.inner_text("#co-sheet-overlay .co-tl")
    check("RDV avec Paul Durand : Salon EnerGaïa" in hist and "Fait : Envoyer la proposition" in hist and "Toit Durand → Négociation" in hist,
          f"company sheet: history of meetings, done tasks and pipeline steps ({hist[:120]!r})")
    page.locator("#co-sheet-overlay .co-row", has_text="Toit Durand").click(); page.wait_for_timeout(200)
    check(page.is_hidden("#co-sheet-overlay") and page.is_visible("#opp-detail-overlay"), "company sheet: a line opens the opportunity")
    page.click(".opp-co-link"); page.wait_for_timeout(100)
    check(page.is_visible("#co-sheet-overlay") and page.inner_text("#co-sheet-title") == "Solaris", "company sheet: reached from the opportunity too")
    page.keyboard.press("Escape"); page.wait_for_timeout(100)
    page.keyboard.press("Control+k"); page.keyboard.type("solaris")
    page.locator("#gs-results .gs-item", has_text="Solaris").filter(has_text="contact").first.click(); page.wait_for_timeout(100)
    check(page.is_visible("#co-sheet-overlay"), "company sheet: reached from the global search")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Today: the next 7 days and deals that stopped moving, under the day's tasks
    old_ms = int((datetime.datetime.now() - datetime.timedelta(days=20)).timestamp() * 1000)
    in3 = (TODAY + datetime.timedelta(days=3)).isoformat()
    td_seed = {"todo_todo-widget-tasks": json.dumps([task("d1", "Préparer le comité", due=in3)]),
               "todo_todo-widget-opportunities": json.dumps([{"id": "do1", "kind": "installation", "company": "Solaris", "centrale": "Toit Durand", "title": "Toit Durand",
                                                             "stage": "pricing_indicatif", "createdAt": old_ms, "updatedAt": old_ms, "deleted": False},
                                                            {"id": "do2", "kind": "installation", "company": "Solaris", "centrale": "Hangar", "title": "Hangar",
                                                             "stage": "pricing_indicatif", "createdAt": 1, "updatedAt": int(datetime.datetime.now().timestamp() * 1000), "deleted": False}])}
    page, errors = new_page(browser, td_seed, viewport={"width": 1300, "height": 900})
    page.goto(URL)
    page.wait_for_function("document.getElementById('today-watch').textContent.includes('Toit Durand')", timeout=10000)
    watch = page.inner_text("#today-watch")
    check("Préparer le comité" in watch and "Toit Durand" in watch and "20 j" in watch and "Hangar" not in watch,
          "today: next 7 days listed, and the deal without movement for 20 days (not the active one)")
    page.locator("#today-watch .tw-row", has_text="Toit Durand").click(); page.wait_for_timeout(200)
    check(page.is_visible("#opp-detail-overlay"), "today: a stalled deal opens in one click")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline dashboard: MWc by phase, time per step, deals that drag
    day_ms = 86400000; now_ms = int(datetime.datetime.now().timestamp() * 1000)
    pd_opps = [{"id": "p1", "kind": "installation", "company": "A", "centrale": "Alpha", "title": "Alpha", "stage": "negociation", "kwc": "2000", "createdAt": now_ms - 40 * day_ms,
                "updatedAt": now_ms - 20 * day_ms, "deleted": False, "stageLog": [{"stage": "pricing_indicatif", "at": now_ms - 40 * day_ms}, {"stage": "negociation", "at": now_ms - 20 * day_ms}]},
               {"id": "p2", "kind": "installation", "company": "B", "centrale": "Beta", "title": "Beta", "stage": "pricing_indicatif", "kwc": "500", "createdAt": now_ms - 2 * day_ms,
                "updatedAt": now_ms - 2 * day_ms, "deleted": False, "stageLog": [{"stage": "pricing_indicatif", "at": now_ms - 2 * day_ms}]},
               {"id": "p3", "kind": "installation", "company": "C", "centrale": "Gamma", "title": "Gamma", "stage": "perdu", "kwc": "900", "lostAt": now_ms - day_ms,
                "createdAt": 1, "updatedAt": now_ms - day_ms, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(pd_opps)}, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#pipe-dash-toggle")
    dash = page.inner_text("#pipeline-board")
    check("2,5 MWc" in dash and "sortis en" in dash.lower() and page.locator(".pd-bar-row").count() == 5, "pipeline dashboard: MWc in progress, exits of the year, one bar per phase")
    rows = page.eval_on_selector_all(".pd-trow", "els=>els.map(e=>e.innerText.replace(/\\s+/g,' '))")
    check(any(r.startswith("Pricing indicatif 11 j") for r in rows) and any(r.startswith("Négociation 20 j") for r in rows), f"pipeline dashboard: average days per step from the history ({rows})")
    check("Alpha" in page.inner_text(".pd-card:has-text('Dossiers qui traînent')"), "pipeline dashboard: the deal stuck for 20 days is listed")
    page.click("#pipe-dash-toggle")
    check(page.locator(".pipe-col").count() >= 5, "pipeline dashboard: one click back to the board")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Accessibility: keyboard, names for screen readers
    page, errors = new_page(browser, {"todo_todo-widget-tasks": json.dumps([task("k1", "Au clavier")])}, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.click('[data-tab="tasks"]')
    chk = page.locator("#task-list .task-line:has-text('Au clavier') .task-check").first
    check(chk.get_attribute("role") == "checkbox" and chk.get_attribute("tabindex") == "0" and chk.get_attribute("aria-label") == "Valider", "a11y: the check circle is a named, focusable checkbox")
    chk.focus(); page.keyboard.press("Space")
    page.wait_for_function(f"({PEEK}('todo-widget-tasks')||'').includes('\"archived\":true')", timeout=10000)
    check(stored(page, "todo-widget-tasks")[0]["archived"] is True, "a11y: a task can be validated with the keyboard")
    unnamed = page.evaluate("[...document.querySelectorAll('button')].filter(b=>b.offsetParent && !b.textContent.trim() && !b.getAttribute('aria-label')).length")
    check(unnamed == 0, f"a11y: every visible icon button has a name ({unnamed} without)")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Excel bridge: never from a phone, even on the full version with a key stored
    ph_seed = {"todo_bridge-config": json.dumps({"port": 8765, "token": "k1"}),
               "todo_todo-widget-opportunities": json.dumps([{"id": "x1", "kind": "installation", "company": "Solaris", "centrale": "Toit", "title": "Toit",
                                                             "stage": "negociation", "createdAt": 1, "updatedAt": 1, "deleted": False}])}
    calls = []
    page, errors = new_page(browser, ph_seed, viewport={"width": 390, "height": 800}, is_mobile=True, has_touch=True,
                            user_agent="Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/18.0 Mobile/15E148 Safari/604.1")
    page.route("http://127.0.0.1:8765/**", lambda r: (calls.append(r.request.url), r.fulfill(status=200, body="{}")))
    page.goto(URL); page.wait_for_timeout(500)
    page.click('[data-tab="pipeline"]'); page.click(".pipe-card:has-text('Toit')")
    acts = page.inner_text("#opp-detail-actions")
    check("Envoyer" not in acts and not calls, f"bridge: no « send to Excel » on a phone, not even on the full version, and no call to the bridge ({calls})")
    page.keyboard.press("Escape"); page.click('[data-tab="lists"]')
    check(page.is_visible("#bridge-phone-note") and not page.is_visible("#bridge-token"), "bridge: Réf. & data says the bridge is for the computer only")
    page.close()

    # --- Windows: « Modifier la fiche » and the closing cross at the top right
    win_opp = [{"id": "w1", "kind": "installation", "company": "SOLARIS", "producteur": "SOLARIS", "centrale": "Toit Durand", "title": "Toit Durand",
                "stage": "collecte", "kwc": "299", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(win_opp)}, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click(".pipe-card")
    box = page.locator("#opp-detail-overlay .modal-box").bounding_box(); eb = page.locator("#opp-detail-edit").bounding_box(); xb = page.locator("#opp-detail-overlay .modal-x").bounding_box()
    check(eb["y"] - box["y"] < 60 and box["x"] + box["width"] - (xb["x"] + xb["width"]) < 30 and xb["x"] > eb["x"] + eb["width"] - 1 and xb["y"] - box["y"] < 40,
          "windows: « Modifier la fiche » at the top, the cross in the top right corner")
    check("Modifier la fiche" not in page.inner_text("#opp-detail-actions"), "windows: no second « Modifier la fiche » at the bottom")
    page.click("#opp-detail-overlay .modal-x")
    check(page.is_hidden("#opp-detail-overlay"), "windows: the cross closes the sheet")
    page.click(".pipe-card"); page.click("#opp-detail-edit")
    check(page.is_visible("#pq-overlay"), "windows: « Modifier la fiche » opens the form")
    page.click("#pq-overlay .modal-x")
    check(page.is_hidden("#pq-overlay") and page.locator(".modal-overlay:not(.hidden)").count() == 0, "windows: the cross closes the form too (every window has one)")
    check(page.evaluate("[...document.querySelectorAll('.modal-overlay .modal-box')].every(b=>b.querySelector('.modal-x, .co-sheet-x'))"), "windows: every window has its cross")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- A pricing request with two installations: tabs and a side-by-side check in each sheet
    two = [{"id": "rq", "kind": "request", "company": "GPX", "title": "Demande de pricing", "createdAt": 1, "updatedAt": 1, "deleted": False},
           {"id": "t1", "kind": "installation", "company": "GPX", "centrale": "ALPHA PV1", "title": "ALPHA PV1", "kwc": "299", "mwh": "358", "mesPrev": "20/07/2027",
            "dept": "49 - Maine-et-Loire", "commune": "VILLE-SUR-EXEMPLE", "requestId": "rq", "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False},
           {"id": "t2", "kind": "installation", "company": "GPX", "centrale": "BETA PV1", "title": "BETA PV1", "kwc": "211.14", "dept": "49 - Maine-et-Loire",
            "requestId": "rq", "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(two)}, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.locator(".pipe-req-name", has_text="ALPHA PV1").click()
    tabs = page.eval_on_selector_all("#opp-detail-body .req-tab", "els=>els.map(e=>e.innerText.replace(/\\s+/g,' ').trim())")
    check(tabs == ["Installation 1 ALPHA PV1 299 kWc · 358 MWh · MES 20/07/2027", "Installation 2 BETA PV1 211 kWc"] and "ALPHA PV1" in page.inner_text("#opp-detail-title"),
          f"request: both installations shown as tabs in the sheet ({tabs})")
    table = page.inner_text("#opp-detail-body .req-compare")
    cellv = lambda oid, label: page.input_value(f"#opp-detail-body .req-in[data-id='{oid}'][data-l=\"{label}\"]")
    check(cellv("t1", "P50 annuel (indicatif) (MWh/an)") == "358" and "3 cases à compléter" in table
          and page.locator("#opp-detail-body .req-table td.empty").count() == 3, "request: side-by-side check, empty cells to complete flagged")
    labels = page.eval_on_selector_all("#opp-detail-body .req-table tbody th.req-lab", "els=>els.map(e=>e.textContent)")
    check(labels[:4] == ["Nom Installation", "Nom SPV", "SIRET SPV", "Adresse Installation"], f"request: lines in the order of the characteristics file ({labels[:4]})")
    # a cell edited in the table goes to the installation (an app field, or the line of the producer's file)
    page.fill("#opp-detail-body .req-in[data-id='t2'][data-l='Commune']", "VILLERS"); page.press("#opp-detail-body .req-in[data-id='t2'][data-l='Commune']", "Enter")
    page.fill("#opp-detail-body .req-in[data-id='t2'][data-l='Latitude']", "47.21"); page.press("#opp-detail-body .req-in[data-id='t2'][data-l='Latitude']", "Enter")
    page.fill("#opp-detail-body .req-in[data-id='t2'][data-l='Département']", "Lot"); page.press("#opp-detail-body .req-in[data-id='t2'][data-l='Département']", "Enter")
    page.wait_for_timeout(200)
    t2 = next(o for o in stored(page, "todo-widget-opportunities") if o["id"] == "t2")
    check(t2.get("commune") == "VILLERS" and t2.get("caract", {}).get("Latitude") == "47.21" and t2.get("dept") == "46 - Lot",
          f"request: edits in the table are saved on the installation ({t2.get('commune')}, {t2.get('caract')}, {t2.get('dept')})")
    page.click("#opp-detail-body tbody[data-sec='Localisation'] .req-sec-h")
    check(not page.is_visible("#opp-detail-body .req-in[data-id='t2'][data-l='Commune']") and page.is_visible("#opp-detail-body .req-in[data-id='t2'][data-l='Nom SPV']"),
          "request: a sub-section folds on its own")
    page.click("#opp-detail-body tbody[data-sec='Localisation'] .req-sec-h")
    check("caractéristiques centrales" in table.lower() and page.locator("#opp-detail-body .req-table td.none").count() > 10,
          "request: the table also lists every line of the plant characteristics file")
    page.click("#opp-detail-body .req-compare-h")
    check(not page.is_visible("#opp-detail-body .req-table"), "request: the side-by-side table folds with a click on its title")
    page.locator("#opp-detail-body .req-tab", has_text="BETA PV1").click()
    check(not page.is_visible("#opp-detail-body .req-table"), "request: folded stays folded on the next sheet")
    page.click("#opp-detail-body .req-compare-h")
    page.locator("#opp-detail-body .req-tab", has_text="ALPHA PV1").click()
    page.locator("#opp-detail-body .req-tab", has_text="BETA PV1").click()
    check("BETA PV1" in page.inner_text("#opp-detail-title") and page.locator("#opp-detail-body .req-tab.on").inner_text().startswith("Installation 2"), "request: a tab opens the other installation")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Duplicate contacts: offered for merging, links follow the kept card
    dup_cts = [{"id": "d1", "firstName": "Louis", "lastName": "Catroux", "role": "", "company": "GPX", "email": "l.catroux@exemple.fr", "phone1": "07 00 00 96 58",
                "phone2": "", "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False},
               {"id": "d2", "firstName": "Louise", "lastName": "Catroux", "role": "Chargée d'affaires", "company": "GPX", "email": "L.Catroux@exemple.fr",
                "phone1": "07 00 00 96 58", "phone2": "02 00 00 00 11", "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 5, "deleted": False},
               {"id": "d3", "firstName": "Paul", "lastName": "Autre", "role": "", "company": "GPX", "email": "p.autre@exemple.fr", "phone1": "", "phone2": "",
                "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    dup_opp = [{"id": "do", "kind": "installation", "company": "GPX", "centrale": "Toit", "title": "Toit", "prenom": "Louise", "nom": "CATROUX", "mail": "l.catroux@exemple.fr",
                "otherContactIds": ["d1", "d3"], "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    dup_mt = [{"id": "m1", "contactId": "d1", "label": "Salon", "date": "2026-10-01", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(dup_cts), "todo_todo-widget-opportunities": json.dumps(dup_opp),
                                      "todo_todo-widget-meetings": json.dumps(dup_mt)}, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click(".pipe-card")
    check(page.is_visible("#opp-detail-body .os-dup") and "même personne" in page.inner_text("#opp-detail-body .os-dup"), "duplicates: the sheet offers to merge the same person")
    page.click("#opp-detail-body .os-dup button")
    rows = page.inner_text("#ct-merge-body")
    check(page.is_visible("#ct-merge-overlay") and "même mail" in rows and "Louise" in rows and "Louis" in rows and "Paul" not in rows, "duplicates: the merge window shows the two cards and why")
    check(page.is_checked("input[name='ctm-0-firstName'][value='Louise']"), "duplicates: the most recent value is chosen by default")
    page.click("#ct-merge-body .ctm-merge")
    live = [c for c in stored(page, "todo-widget-contacts") if not c.get("deleted")]
    o = stored(page, "todo-widget-opportunities")[0]
    mt = stored(page, "todo-widget-meetings")[0]
    check(len(live) == 2 and any(c["firstName"] == "Louise" and c["phone2"] == "02 00 00 00 11" for c in live) and o["otherContactIds"] == ["d3"] and mt["contactId"] == "d2",
          f"duplicates: one card kept, the opportunity and the meeting follow it ({[c['firstName'] for c in live]}, {o.get('otherContactIds')}, {mt['contactId']})")
    page.click("#ct-merge-overlay .modal-x")
    check(not page.is_visible("#opp-detail-body .os-dup"), "duplicates: the sheet no longer warns once merged")
    page.click("#opp-detail-overlay .modal-x"); page.click('[data-tab="contacts"]')
    check(page.is_hidden("#ct-dupes-banner"), "duplicates: no banner left once merged")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Contacts page: list view by default, pencil to edit, top bars always in view
    many = [{"id": f"k{i}", "firstName": f"P{i}", "lastName": f"Nom{i:02d}", "role": "", "company": "Gamma", "email": "", "phone1": "", "phone2": "",
             "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False} for i in range(40)]
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(many)}, viewport={"width": 1300, "height": 700})
    page.goto(URL); page.evaluate("localStorage.removeItem('contacts-view')"); page.reload(); page.click('[data-tab="contacts"]')
    check(page.locator('#contacts-view-switch [data-view="liste"].active').count() == 1 and page.locator("#contacts-list .company-group-header").count() == 0,
          "contacts: list view by default")
    page.locator("#contacts-list .contact-row", has_text="Nom05").locator(".edit-btn").click()
    check(page.is_visible("#contact-form") and page.input_value("#ct-last") == "Nom05", "contacts: the pencil opens the contact to edit")
    page.click("#contact-form >> text=Annuler") if page.locator("#contact-form >> text=Annuler").count() else None
    page.evaluate("document.getElementById('tab-contacts').scrollTop=1500"); page.wait_for_timeout(100)
    sb = page.locator("#contact-search").bounding_box()
    check(sb and 0 <= sb["y"] < 80, "contacts: search and buttons stay in view while scrolling")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- « Caractéristiques Centrales » file generated with every line of the template, and a pricing mail with the figures
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Données"
    for r in [["", "Centrale 1", "Centrale 2"], ["Nom Installation", "Les Pins", "Le Lac"], ["Nom SPV", "PINS SAS", "LAC SAS"], ["Adresse Installation", "LD Les Pins, 11380 Villeneuve", "LD Le Lac, 47110 Exemple"],
              ["Département", "Aude (11)", "Lot-et-Garonne (47)"], ["Lattitude", "43.38", "44.39"], ["Longitude", "2.43", "0.55"], ["Nature Projet", "Centrale au sol", "Toiture"],
              ["Première injection ", "Mars 2027", "Avril 2027"], ["Date de mise en service (effective ou prévisionnelle)", "Avril 2027", "Mai 2027"],
              ["Puissance raccordement (kVA)", 4200, 3170], ["Puissance installée (kWc)", 4986.93, 4790.89], ["P50 annuel (indicatif) (MWh/an)", 6624.6, 6793],
              ["P90 annuel (indicatif) (MWh/an)", 6203, 6315], ["Champ 1 : Orientation (° par rapport au sud)", "0° (plein sud)", "0° (plein sud)"], ["Champ 1 : Inclinaison (°)", 20, 20]]:
        ws.append(r)
    src = pathlib.Path(tempfile.mkdtemp()) / "carac_src.xlsx"; wb.save(src)
    page, errors = new_page(browser, viewport={"width": 1300, "height": 900}, accept_downloads=True)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "EXEMPLE ENERGIES"); page.click("#pq-next-btn")
    page.set_input_files("#pq-file", str(src)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    page.click("#pq-save-btn"); page.wait_for_timeout(200)
    if page.is_visible("#pq-count-ok"): page.check("#pq-count-ok"); page.click("#pq-save-btn"); page.wait_for_timeout(300)
    page.keyboard.press("Escape"); page.wait_for_timeout(100)
    for ov in ("#opp-detail-overlay", "#pq-overlay"):
        if page.is_visible(ov + " .modal-x"): page.click(ov + " .modal-x")
    page.locator(".pipe-req-card .pipe-card-title").first.click(); page.wait_for_timeout(200)
    mail = page.input_value("#opp-mail-body")
    check("Production : voir fichier joint." in mail and "Les Pins : mars 2027 → avril 2027 (test 1 mois)" in mail and "Le Lac : avril 2027 → mai 2027 (test 1 mois)" in mail,
          f"pricing mail: figures in the attached file, first injection → MSI from the producer's file ({[l for l in mail.split(chr(10)) if 'Les Pins' in l or 'Le Lac' in l]})")
    with page.expect_download() as dl:
        page.click("#opp-detail-actions >> text=Fichier caractéristiques (.xlsx)")
    got_file = pathlib.Path(tempfile.mkdtemp()) / dl.value.suggested_filename; dl.value.save_as(got_file)
    check(got_file.name.endswith("_Caractéristiques_Centrales_EXEMPLE ENERGIES.xlsx"), f"caract file: named like the usual files ({got_file.name})")
    out = openpyxl.load_workbook(got_file)["Données"]
    rows = {r[0].value: [c.value for c in r[1:]] for r in out.iter_rows(min_row=2)}
    check(out.max_row == 49 and out["B1"].value == "Centrale 1" and out["C1"].value == "Centrale 2" and "Champ 4 : Fixe ou tracker" in rows,
          f"caract file: every line of the template ({out.max_row} rows), one column per plant")
    check(rows["Nom Installation"] == ["Les Pins", "Le Lac"] and rows["Latitude"] == ["43.38", "44.39"] and rows["Puissance raccordement (kVA)"] == [4200, 3170]
          and rows["Première injection"] == ["Mars 2027", "Avril 2027"] and rows["Puissance installée (kWc)"] == [4986.93, 4790.89] and rows["Département"] == ["Aude (11)", "Lot-et-Garonne (47)"]
          and rows["Champ 1 : Inclinaison (°)"] == [20, 20] and rows["Stockage"] == [None, None],
          f"caract file: the app's values and every line of the producer's file ({rows.get('Latitude')}, {rows.get('Première injection')})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Contacts import: a conflict says which field and where each value comes from
    imp_ct = [{"id": "ic1", "firstName": "Anne", "lastName": "Roux", "role": "", "company": "Solaris", "email": "anne@exemple.fr", "phone1": "06 11 22 33 44",
               "phone2": "", "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Prénom", "Nom", "Société", "Email", "Téléphone"]); ws.append(["Anne", "Roux", "Solaris", "anne@exemple.fr", "06 99 88 77 66"])
    imp = pathlib.Path(tempfile.mkdtemp()) / "contacts.xlsx"; wb.save(imp)
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(imp_ct)}, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="contacts"]')
    page.on("dialog", lambda d: d.accept())
    page.set_input_files("#contact-import-input", str(imp)); page.wait_for_timeout(1500)
    page.click('[data-tab="lists"]')
    conf = page.inner_text("#conflicts-list") if page.locator("#conflicts-list").count() else ""
    check("Contact « Anne Roux » — champ « Téléphone 1 »" in conf and "Déjà dans tes contacts" in conf and "Dans le fichier « contacts.xlsx », ligne 2" in conf and "06 99 88 77 66" in conf,
          f"import conflict: the field and where each value comes from ({conf[:160]!r})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Post-it edited on two devices: lines merged, no conflict reported (old ones cleared)
    base_n = [{"id": "scratchpad", "text": "Appeler Paul\nEnvoyer offre", "createdAt": 1, "updatedAt": 10, "deleted": False}]
    loc_n = [{"id": "scratchpad", "text": "Appeler Paul\nEnvoyer offre\nIdée PC", "createdAt": 1, "updatedAt": 20, "deleted": False}]
    files = {"notes.json": [{"id": "scratchpad", "text": "Envoyer offre\nIdée téléphone", "createdAt": 1, "updatedAt": 30, "deleted": False}],
             "tasks.json": [], "crm.json": {"companies": [], "contacts": [], "meetings": [], "opportunities": []}}
    old_conf = [{"id": "cfn", "entityType": "note", "entityId": "scratchpad", "entityName": "Post-it", "field": "text", "fieldLabel": "Texte", "localValue": "a", "remoteValue": "b", "resolvedValue": "a"}]
    page, errors = new_page(browser, dict(SYNC_ON, **{"todo_todo-widget-notes": json.dumps(loc_n), "todo_sync-notes-base": json.dumps(base_n),
                                                     "todo_sync-conflicts": json.dumps(old_conf)}), viewport={"width": 1300, "height": 850})
    mock_github(page, files); page.goto(URL)
    page.wait_for_function("document.getElementById('postit-text').value.includes('Idée téléphone')", timeout=10000)
    txt = page.input_value("#postit-text")
    check(txt == "Envoyer offre\nIdée PC\nIdée téléphone" or txt == "Envoyer offre\nIdée téléphone\nIdée PC",
          f"post-it: both devices' lines kept, the line deleted on the other device removed ({txt!r})")
    check(not [c for c in (stored(page, "sync-conflicts") or []) if c.get("entityType") == "note"], "post-it: no conflict reported, old ones cleared")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Post-it: a checklist by default (tick, struck through, kept in the text); click a line to edit
    page, errors = new_page(browser, {"todo_todo-widget-notes": json.dumps([{"id": "scratchpad", "text": "Appeler Paul\nEnvoyer offre", "createdAt": 1, "updatedAt": 1, "deleted": False}])},
                            viewport={"width": 1300, "height": 850})
    page.goto(URL); page.wait_for_selector("#postit-list .pi-row")
    check(page.locator("#postit-list .pi-row input[type=checkbox]").count() == 2 and not page.is_visible("#postit-text"), "post-it: each line shown with a box to tick")
    page.locator("#postit-list .pi-row", has_text="Appeler Paul").locator("input").check()
    page.wait_for_function(f"({PEEK}('todo-widget-notes')||'').includes('✓ Appeler Paul')", timeout=5000)
    check(page.locator("#postit-list .pi-row.done").count() == 1, "post-it: a ticked line is struck through and saved")
    page.locator("#postit-list .pi-text", has_text="Envoyer offre").click()
    check(page.is_visible("#postit-text") and page.evaluate("document.activeElement.id") == "postit-text", "post-it: clicking a line edits the text")
    page.keyboard.type(" PPA"); page.locator("#postit-bar").click(); page.wait_for_timeout(200)
    check([t.strip() for t in page.eval_on_selector_all("#postit-list .pi-text", "els=>els.map(e=>e.textContent)")] == ["Appeler Paul", "Envoyer offre PPA"],
          "post-it: back to the checklist after editing")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Power, P50 and specific yield: any two give the third; P90 suggested; département from the postcode
    page, errors = new_page(browser, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA CALCUL"); page.click("#pq-next-btn")
    def put(k, v):
        page.fill("#pq-0-" + k, v); page.locator("#pq-0-" + k).dispatch_event("input"); page.locator("#pq-0-" + k).dispatch_event("change")
    is_auto = lambda k: page.evaluate(f"document.querySelector('#pq-0-{k}').closest('.pq-f').classList.contains('auto')")
    put("kwc", "300"); put("productible", "1200")
    check(page.input_value("#pq-0-mwh") == "360" and is_auto("mwh"), f"kWc × productible → P50 computed, in green ({page.input_value('#pq-0-mwh')})")
    put("productible", "1300")
    check(page.input_value("#pq-0-mwh") == "390", "a new productible updates the computed P50")
    put("productible", "")
    check(page.input_value("#pq-0-mwh") == "", "a computed P50 goes away with its source")
    put("mwh", "450")
    check(page.input_value("#pq-0-productible") == "1500" and is_auto("productible"), "kWc and P50 → productible")
    put("productible", "1400")
    check(page.input_value("#pq-0-mwh") == "450" and page.input_value("#pq-0-productible") == "1400", "typed values are never overwritten")
    page.click("#pq-form [data-p90]")
    check(page.input_value("#pq-0-p90") == "414", f"P90 suggested from the P50 (92 % in solar) ({page.input_value('#pq-0-p90')})")
    put("cp", "46100")
    check(page.input_value("#pq-0-dept") == "46 - Lot", f"département from the postcode ({page.input_value('#pq-0-dept')})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Business code (GENE01-D1, agrégation): month only when the start is not 1 January / the end not 31 December
    page, errors = new_page(browser); page.add_init_script("window.TODO_TEST_HOOKS=true;"); page.goto(URL)
    codes = page.evaluate("""[
      {producteur:'SOLARIS', centrale:'Toit Durand', debut:'01/01/2027', fin:'31/12/2029'},
      {producteur:'SOLARIS', centrale:'Toit Durand', debut:'01/06/2027', fin:'31/08/2030'},
      {company:'GAEC Martin', centrale:'Hangar 2', mesPrev:'15/02/2028', duree:'3'},
      {producteur:'SOLARIS', centrale:'CS du Moulin'},
      {producteur:'SOLARIS', centrale:'Ombrière', mesPrev:'2027', fin:'31/12/2030'}].map(o=>window.__todo.codeAffaire(Object.assign({kind:'installation'},o)))""")
    short = page.evaluate("[window.__todo.codeAffaire({kind:'installation', producteur:'P', centrale:'C', debut:'juil-27', fin:'31/12/2030'})]")
    check(short == ["27(07)30 – P_C – A"], f"Excel's short dates (« juil-27 ») are read ({short})")
    check(codes == ["2729 – SOLARIS_Toit Durand – A", "27(06)30(08) – SOLARIS_Toit Durand – A", "28(02)31(02) – GAEC Martin_Hangar 2 – A", "[AA][BB] – SOLARIS_CS du Moulin – A", "2730 – SOLARIS_Ombrière – A"],
          f"business code GENE01-D1 (agrégation): AA(MM)BB(MM) - Producteur_Centrale - A ({codes})")
    # the pricing mail, on the model of the last one sent (invented plants)
    ml = page.evaluate("""(()=>{ const base={ kind:'installation', company:'EXEMPLE ENERGIE', producteur:'EXEMPLE ENERGIE', contrat:'AO CRE CR', fin:'31/12/2030' };
      const L=[Object.assign({}, base, { centrale:'Les Pins', install:'PV', debut:'01/03/2027', mesPrev:'01/04/2027' }),
               Object.assign({}, base, { centrale:'Le Lac', install:'Eolien', debut:'01/12/2027', mesPrev:'12/01/2028', fin:'31/12/2031' }),
               Object.assign({}, base, { centrale:'Les Prés', install:'PV', debut:'01/07/2027', mesPrev:'01/08/2027' })];
      return window.__todo.pricingMail({ company:'EXEMPLE ENERGIE' }, L); })()""")
    b = ml["body"]
    check(ml["subject"] == "Demande de prix agrégation – EXEMPLE ENERGIE – 2 PV + 1 éolien CR", f"pricing mail: subject like the model ({ml['subject']})")
    check("pour EXEMPLE ENERGIE : 2 projets PV greenfield et 1 projet éolien greenfield, tous lauréats AO CRE et bénéficiant d’un contrat de complément de rémunération." in b
          and "- 27(03)30 – EXEMPLE ENERGIE_Les Pins – A" in b and "(infos en jaune extrapolées)" in b, "pricing mail: intro and business codes like the model")
    order = [l.strip() for l in b.split("\n") if "(test " in l]
    check(order == ["- Les Pins : mars 2027 → avril 2027 (test 1 mois)", "- Les Prés : juillet 2027 → août 2027 (test 1 mois)", "- Le Lac : décembre 2027 → janvier 2028 (test 6 semaines)"],
          f"pricing mail: first injection → MSI per plant, in order, with the test length ({order})")
    check("- Durée : 2 pricing demandés, jusqu’au 31/12/2030 et jusqu’au 31/12/2031." in b and "- Contexte du marché : AO CRE (CR)" in b
          and "   - Prix phase de Test : SPOT" in b and "   - Période de test : frais en €/MWh" in b and "- GO : sans objet" in b and "- GAPD : coût et montant à indiquer." in b
          and "- Éolien : bridages" in b and "Date de remise de l’offre : [À CONFIRMER" in b and b.endswith("Merci d’avance !"), "pricing mail: options, test phase and conditions like the model")
    page.close()

    # --- Contract dates: start + duration → end, start + end → duration, one-click choices
    page, errors = new_page(browser, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA DATES"); page.click("#pq-next-btn")
    page.click("#pq-fold")
    check(not page.is_visible("#pq-mail") and page.get_attribute("#pq-fold", "aria-expanded") == "false", "form: the mail / file column folds away")
    page.click("#pq-fold"); check(page.is_visible("#pq-mail"), "form: and opens again")
    y = datetime.date.today().year
    page.click(f"#pq-form [data-dset='fin'][data-v='31/12/{y + 2}']")
    check(page.input_value("#pq-0-fin") == f"31/12/{y + 2}", "dates: end on 31/12 of next year, the year after or in 3 years in one click")
    page.fill("#pq-0-fin", ""); page.locator("#pq-0-fin").dispatch_event("change")
    def setv(k, v, i=0):
        page.fill(f"#pq-{i}-{k}", v); page.locator(f"#pq-{i}-{k}").dispatch_event("change")
    setv("mesPrev", "20/07/2027")
    page.click("#pq-form button[data-dset='debut']:has-text('MES')")
    check(page.locator("#pq-form [data-dset='duree']").count() == 0, "dates: no duration chips (the duration never sets the end)")
    setv("duree", "3")
    check(page.input_value("#pq-0-debut") == "20/07/2027" and page.input_value("#pq-0-fin") == "", "dates: a duration typed alone never changes the end date")
    setv("fin", "19/07/2030")
    check(page.input_value("#pq-0-duree") == "3", f"dates: start and end give the duration ({page.input_value('#pq-0-duree')})")
    setv("fin", "31/12/2031")
    check(page.input_value("#pq-0-duree") == "4,5", f"dates: a new end date recomputes the duration ({page.input_value('#pq-0-duree')})")
    setv("debut", "janvier 2028")
    check(page.input_value("#pq-0-debut") == "01/01/2028" and page.input_value("#pq-0-duree") == "4" and page.input_value("#pq-0-fin") == "31/12/2031",
          f"dates: « janvier 2028 » written out, the end stays, the duration follows ({page.input_value('#pq-0-duree')})")
    # several plants: pushing a date pushes that date only; each plant at its own commissioning date
    page.click("#pq-add-inst"); page.click("#pq-add-inst")
    setv("mesPrev", "01/03/2028", 1)
    check(page.locator("#pq-form button[data-dset='debut']:text-is('MES')").count() == 0, "dates: with several installations, only « MES toutes » (no single « MES »)")
    page.click("#pq-form button[data-toall='fin'][data-s='0']")
    check(page.is_visible("#choice-overlay") and "Date fin" in page.inner_text("#choice-text"), "« → toutes »: a round button, confirmed before applying")
    page.click("#choice-btns >> text=Appliquer à toutes")
    check([page.input_value(f"#pq-{i}-fin") for i in range(3)] == ["31/12/2031"] * 3 and page.input_value("#pq-1-debut") == "" and page.input_value("#pq-2-debut") == "",
          "« → toutes » on the end date pushes the end date only (never the start)")
    check(page.evaluate("""(()=>{ const b=document.querySelector("#pq-form button[data-dall='mes'][data-s='0']"), f=b.closest('.pq-f'), r=b.getBoundingClientRect(), p=f.getBoundingClientRect();
          return !!f.querySelector('#pq-0-debut') && r.right<=p.right+1 && b.scrollWidth<=b.clientWidth+1; })()"""), "dates: « MES toutes » under the start date, fully visible")
    page.click("#pq-form button[data-dall='mes'][data-s='0']")
    check(page.is_visible("#choice-overlay"), "dates: asks before replacing start dates already filled")
    page.click("#choice-btns >> text=Remplacer")
    got = [[page.input_value(f"#pq-{i}-{k}") for k in ("debut", "duree", "fin")] for i in range(3)]
    check(got[0] == ["20/07/2027", "4,5", "31/12/2031"] and got[1] == ["01/03/2028", "4", "31/12/2031"] and got[2] == ["", "", "31/12/2031"],
          f"dates: each plant at its own commissioning date, the ends unchanged, the durations follow ({got})")
    page.select_option("#pq-0-contrat", "AOS CRE"); page.click("#pq-form button[data-toall='contrat'][data-s='0']"); page.click("#choice-btns >> text=Appliquer à toutes")
    check([page.input_value(f"#pq-{i}-contrat") for i in range(3)] == ["AOS CRE"] * 3, "« → toutes » puts any field of an installation on all the others")
    cls = lambda k, i=0: page.get_attribute(f"#pq-form button[data-toall='{k}'][data-s='{i}']", "class")
    check("off" in cls("particularites") and "Déjà" in page.get_attribute("#pq-form button[data-toall='contrat'][data-s='0']", "aria-label")
          and page.locator("#pq-form button[data-toall='centrale']").count() == 0 and page.locator("#pq-form button[data-toall='kwc']").count() == 0,
          "« → toutes »: hidden while empty, ✓ when all the same, never on a plant's name or power")
    page.fill("#pq-0-particularites", "Toiture"); page.locator("#pq-0-particularites").dispatch_event("change")
    page.select_option("#pq-0-install", "PV"); page.wait_for_timeout(100)
    check("go" in cls("particularites"), "« → toutes » stands out when the others differ")
    page.click("#pq-form button[data-toall='particularites'][data-s='0']"); page.click("#choice-btns >> text=Appliquer à toutes")
    check([page.input_value(f"#pq-{i}-particularites") for i in range(3)] == ["Toiture"] * 3, "« → toutes » for Particularités too")
    check(page.locator("#pq-form .pq-inst.many .pq-inst-n").count() == 3, "each installation in its own numbered card")
    page.select_option("#pq-0-contrat", ""); page.locator("#pq-0-contrat").dispatch_event("change"); page.wait_for_timeout(100)
    check("clear" in cls("contrat"), "« → toutes » on an empty field when the others are filled")
    page.click("#pq-form button[data-toall='contrat'][data-s='0']")
    check("Vider" in page.inner_text("#choice-btns"), "pushing an empty field asks before emptying it everywhere")
    page.click("#choice-btns >> text=Vider partout")
    check([page.input_value(f"#pq-{i}-contrat") for i in range(3)] == [""] * 3, "« → toutes » with no value empties the field everywhere")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Stages in batch: all the installations of a request move together, or one by one on request
    bt = [{"id": "rqb", "kind": "request", "company": "BATCH", "title": "Demande de pricing", "createdAt": 1, "updatedAt": 1, "deleted": False}] + \
         [{"id": f"b{i}", "kind": "installation", "company": "BATCH", "centrale": f"Centrale {i}", "title": f"Centrale {i}", "requestId": "rqb", "stage": "collecte",
           "createdAt": 1, "updatedAt": 1, "deleted": False} for i in range(1, 4)]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(bt)}, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.locator(".pipe-req-name", has_text="Centrale 1").click()
    check(page.locator("#opp-detail-body .stage-scope button.on").inner_text() == "Toutes les installations (3)", "batch: by default a stage applies to all the installations")
    check(page.inner_text("#opp-detail-sum") == "3 installations", "sheet: the header sums up the opportunity (installations · power · technology · volume)")
    page.click("#opp-detail-body .opp-stage-step:has-text('Demande de prix')"); page.wait_for_timeout(200)
    st_ = {o["id"]: o.get("stage") for o in stored(page, "todo-widget-opportunities")}
    check(st_["b1"] == st_["b2"] == st_["b3"] == "demande_prix", f"batch: the three installations moved together ({st_})")
    page.click("#opp-detail-body .stage-scope button:has-text('Seulement celle-ci')")
    page.click("#opp-detail-body .opp-stage-step:has-text('Pricing indicatif')"); page.wait_for_timeout(200)
    st_ = {o["id"]: o.get("stage") for o in stored(page, "todo-widget-opportunities")}
    check(st_["b1"] == "pricing_indicatif" and st_["b2"] == st_["b3"] == "demande_prix", "batch: one by one is still possible")
    page.click(".opp-back")
    check("1" in page.inner_text("#opp-detail-body .opp-stage-step.current:has-text('Pricing indicatif') .stage-n") and page.locator("#opp-detail-body .req-inst").count() == 3,
          "request sheet: the same step track, with the number of installations at each step")
    page.click("#opp-detail-body .opp-stage-step:has-text('Négociation')"); page.wait_for_timeout(200)
    check(page.inner_text("#opp-detail-edit") == "Modifier les fiches" and page.locator("#opp-detail-body .req-inst-ed").count() == 3, "request sheet: edit buttons in the global view")
    st_ = {o["id"]: o.get("stage") for o in stored(page, "todo-widget-opportunities")}
    check(st_["b1"] == st_["b2"] == st_["b3"] == "negociation", "batch: from the request sheet, all installations to one stage")
    page.locator("#opp-detail-body .req-inst", has_text="Centrale 2").locator(".req-inst-ed").click()
    check(page.is_visible("#pq-overlay") and page.input_value("#pq-0-centrale") == "Centrale 2" and page.locator("#pq-form .pq-inst").count() == 3,
          "request sheet: the pencil opens the form with that installation first, the others below")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Copy for the Pipe: one line per installation or one line for all; « → toutes » asks before replacing
    cp_opps = [{"id": "rqc", "kind": "request", "company": "COPY", "title": "Demande de pricing", "createdAt": 1, "updatedAt": 1, "deleted": False}] + \
              [{"id": f"c{i}", "kind": "installation", "company": "COPY", "producteur": "COPY", "centrale": f"Toit {i}", "title": f"Toit {i}", "kwc": str(100 * i), "mwh": str(120 * i),
                "contrat": "AO CRE" if i == 2 else "", "requestId": "rqc", "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False} for i in (1, 2)]
    ctx = browser.new_context(viewport={"width": 1300, "height": 900}, permissions=["clipboard-read", "clipboard-write"])
    page = ctx.new_page(); errors = []; page.on("pageerror", lambda e: errors.append(str(e)))
    page.route("**/fonts.g*/**", lambda r: r.abort())
    page.add_init_script("if(!sessionStorage.getItem('seeded')){ localStorage.setItem('todo_todo-widget-opportunities', " + json.dumps(json.dumps(cp_opps)) + "); sessionStorage.setItem('seeded','1'); }")
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.locator(".pipe-req-card .pipe-card-title").first.click()
    page.click("#opp-detail-actions >> text=Copier lignes Pipe")
    check(page.is_visible("#choice-overlay") and "2 installations" in page.inner_text("#choice-text"), "pipe copy: asks one line per installation or one for all")
    page.click("#choice-btns >> text=Une seule ligne pour toutes"); page.wait_for_timeout(200)
    clip = page.evaluate("navigator.clipboard.readText()")
    check(clip.count("\n") == 0 and "Toit 1 + Toit 2" in clip and "\t300\t" in clip.replace(",", "."), f"pipe copy: one line, names joined, powers added ({clip[:80]!r})")
    page.click("#opp-detail-actions >> text=Copier lignes Pipe"); page.click("#choice-btns >> text=Une ligne par installation"); page.wait_for_timeout(200)
    check(page.evaluate("navigator.clipboard.readText()").count("\n") == 1, "pipe copy: or one line each")
    page.click("#opp-detail-edit")
    page.select_option("#pq-0-contrat", "PPA Spot"); page.click("#pq-form button[data-toall='contrat'][data-s='0']")
    check(page.is_visible("#choice-overlay") and "AO CRE" in page.inner_text("#choice-text"), "« → toutes »: asks before replacing a value already filled")
    page.click("#choice-btns >> text=Seulement les champs vides")
    check(page.input_value("#pq-1-contrat") == "AO CRE", "« → toutes »: « only empty fields » keeps the existing value")
    page.click("#pq-form button[data-toall='contrat'][data-s='0']"); page.click("#choice-btns >> text=Remplacer partout")
    check(page.input_value("#pq-1-contrat") == "PPA Spot", "« → toutes »: « replace everywhere » replaces it")
    check(not errors, f"no page errors ({errors})")
    ctx.close()

    # --- Undo / redo of actions: several steps back with Ctrl+Z, Ctrl+Shift+Z to redo
    page, errors = new_page(browser, {"todo_todo-widget-tasks": json.dumps([task("u1", "Première"), task("u2", "Deuxième")])}, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.click('[data-tab="tasks"]'); page.wait_for_timeout(300)
    page.locator("#task-list .task-line", has_text="Première").locator(".task-check").first.click(); page.wait_for_timeout(900)
    page.fill("#quick-add-input", "Troisième"); page.press("#quick-add-input", "Enter"); page.wait_for_timeout(900)
    page.locator("body").click(position={"x": 600, "y": 700}); page.keyboard.press("Control+z"); page.wait_for_timeout(200)
    t = {x["title"]: x for x in stored(page, "todo-widget-tasks")}
    check(t["Troisième"]["deleted"] is True and t["Première"]["archived"] is True and "Annulé" in page.inner_text("#bridge-msg"), "undo: Ctrl+Z cancels the last action (the added task)")
    page.keyboard.press("Control+z"); page.wait_for_timeout(200)
    t = {x["title"]: x for x in stored(page, "todo-widget-tasks")}
    check(t["Première"]["archived"] is False, "undo: a second Ctrl+Z goes one more step back (validation cancelled)")
    page.keyboard.press("Control+Shift+z"); page.wait_for_timeout(200)
    t = {x["title"]: x for x in stored(page, "todo-widget-tasks")}
    check(t["Première"]["archived"] is True and t["Troisième"]["deleted"] is True, "undo: Ctrl+Shift+Z redoes")
    page.focus("#quick-add-input"); page.keyboard.type("abc"); page.keyboard.press("Control+z")
    check(len(stored(page, "todo-widget-tasks")) == 3 and t["Première"]["archived"] is True, "undo: inside a text field, Ctrl+Z only undoes the typing")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Sync pill: state visible everywhere, token problems explained
    page, errors = new_page(browser, viewport={"width": 1300, "height": 800})
    page.goto(URL); page.wait_for_timeout(300)
    check(page.get_attribute("#sync-pill", "data-k") == "off" and "désactivée" in page.inner_text("#sync-pill"), "sync pill: says when sync is off")
    page.click("#sync-pill"); page.wait_for_timeout(200)
    check(page.is_visible("#sync-section") and page.evaluate("document.activeElement.id") == "sync-owner", "sync pill: click opens the sync settings")
    page.close()
    files = {"tasks.json": [], "crm.json": {"companies": [], "contacts": [], "meetings": [], "opportunities": []}, "notes.json": []}
    page, errors = new_page(browser, SYNC_ON, viewport={"width": 1300, "height": 800})
    mock_github(page, files); page.goto(URL)
    page.wait_for_function("document.getElementById('sync-pill').dataset.k==='ok'", timeout=8000)
    check("Synchronisé" in page.inner_text("#sync-pill"), "sync pill: green with the time of the last sync")
    page.close()
    page, errors = new_page(browser, SYNC_ON, viewport={"width": 1300, "height": 800})
    page.route("**/api.github.com/repos/o/r/contents/data/*", lambda r: r.fulfill(status=401, body="{}"))
    page.goto(URL); page.wait_for_function("document.getElementById('sync-pill').dataset.k==='err'", timeout=8000)
    check("Jeton refusé" in page.inner_text("#sync-pill"), "sync pill: a refused token is shown in red, everywhere")
    page.close()
    soon = (TODAY + datetime.timedelta(days=10)).isoformat()
    cfg = {"todo_sync-config": json.dumps({"owner": "o", "repo": "r", "token": "t", "enabled": True, "tokenExpires": soon})}
    page, errors = new_page(browser, cfg, viewport={"width": 390, "height": 800}, is_mobile=True, has_touch=True)
    mock_github(page, files); page.goto(URL + "?lite")
    page.wait_for_function("document.getElementById('sync-pill-lite').dataset.k==='warn'", timeout=10000)
    check(page.is_visible("#sync-pill-lite") and page.get_attribute("#sync-pill-lite", "data-k") == "warn" and "10 j" in page.inner_text("#sync-pill-lite"),
          "sync pill: token expiring in 10 days warned, also on the phone version")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Post-it syncs between devices
    files = {"notes.json": [{"id": "scratchpad", "text": "écrit sur le téléphone", "createdAt": 1, "updatedAt": 5, "deleted": False}]}
    page, errors = new_page(browser, SYNC_ON, viewport={"width": 1300, "height": 850})
    mock_github(page, files)
    page.goto(URL); page.wait_for_timeout(800)
    check(page.input_value("#postit-text") == "écrit sur le téléphone", "post-it pulled from the other device")
    postit_edit(page); page.keyboard.press("End"); page.keyboard.type(" + PC")
    page.wait_for_timeout(4500)
    check(files["notes.json"][0]["text"] == "écrit sur le téléphone + PC", "post-it edit pushed")
    page.close()

    # --- Sync: a data file over 1 MB is read through the raw media type
    big = [{"id": "c%d" % i, "firstName": "P%d" % i, "lastName": "N", "company": "", "email": "", "phone1": "", "phone2": "", "role": "",
            "relevance": 0, "comment": "x" * 200, "createdAt": 1, "updatedAt": 1, "deleted": False} for i in range(300)]
    files = {"crm.json": {"companies": [], "contacts": big, "meetings": [], "opportunities": []}, "tasks.json": [], "notes.json": []}
    page, errors = new_page(browser, SYNC_ON)
    mock_github(page, files, large=("crm.json",))
    page.goto(URL); page.wait_for_timeout(1500)
    check(len(stored(page, "todo-widget-contacts") or []) == 300, "sync reads a data file over 1 MB")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Contacts on a phone: no horizontal overflow, delete available in the card
    contacts = [{"id": "c1", "firstName": "Julien", "lastName": "Marchand", "role": "Directeur Commercial", "company": "Alpha Energy",
                 "email": "j@alpha.fr", "phone1": "0612345678", "phone2": "", "relevance": 4, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(contacts)},
                            viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    page.goto(URL)
    check(not page.is_visible("#postit"), "post-it minimised by default on phone")
    page.tap('[data-tab="contacts"]')
    check(page.evaluate("document.documentElement.scrollWidth") <= 390, "contacts fit the phone screen")
    page.tap(".contact-main")
    check(page.locator(".contact-detail-links button", has_text="Supprimer").count() == 1, "delete reachable from the contact card")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Contacts: alphabetical by default, sortable columns, company info on demand, edit keeps scroll
    people = [("Zoé", "Martin", "Beta", 2), ("Adam", "Durand", "Beta", 5), ("Léa", "Bernard", "Alpha", 3), ("Hugo", "Petit", "", 1)]
    people += [(f"P{i}", f"Nom{i:02d}", "Gamma", 1) for i in range(40)]
    cts = [{"id": f"c{i}", "firstName": f, "lastName": l, "role": "", "company": co, "email": "", "phone1": "", "phone2": "",
            "relevance": r, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False} for i, (f, l, co, r) in enumerate(people)]
    cos = [{"id": "co1", "name": "Alpha", "type": "Client", "whatTheyDo": "Producteur hors France", "marketScope": ["EU"], "website": "",
            "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(cts), "todo_todo-widget-companies": json.dumps(cos)},
                            viewport={"width": 1300, "height": 700})
    page.goto(URL); page.click('[data-tab="contacts"]')
    names = lambda: page.eval_on_selector_all("#contacts-list .contact-name > span:first-child", "els=>els.map(e=>e.textContent)")
    check(names()[:3] == ["Léa Bernard", "Adam Durand", "Zoé Martin"], f"contacts alphabetical by default {names()[:3]}")
    check("Producteur hors France" not in page.inner_text("#contacts-list"), "company info hidden by default")
    page.locator(".company-group-header", has_text="Alpha").locator(".co-info-btn").click()
    check("Producteur hors France" in page.inner_text("#contacts-list"), "company info shown with the (i) button")
    page.click('.contacts-header-row .sort[data-sort="rel"]')
    # stars are the client's: Beta starts at 5 (best of its contacts), Alpha at 3, Gamma at 1
    check(names()[:3] == ["Adam Durand", "Zoé Martin", "Léa Bernard"], f"sort by Note puts the best-rated clients first ({names()[:3]})")
    rel = {c["name"]: c.get("relevance") for c in stored(page, "todo-widget-companies")}
    check(rel.get("Alpha") == 3 and not page.is_visible("#ct-relevance-picker"), f"client note starts from its best contact, none on the contact form ({rel})")
    page.locator(".company-group-header", has_text="Alpha").locator(".co-stars .co-star").nth(4).click()
    rel = {c["name"]: c.get("relevance") for c in stored(page, "todo-widget-companies")}
    check(rel.get("Alpha") == 5, "clicking the stars on the client line rates the client")
    page.click('#contacts-view-switch [data-view="liste"]')
    page.click('.contacts-header-row .sort[data-sort="name"]'); page.click('.contacts-header-row .sort[data-sort="name"]')
    check(names()[0] == "Zoé Martin", f"clicking Contact twice sorts Z→A {names()[:2]}")
    page.click('.contacts-header-row .sort[data-sort="name"]')
    check(names()[:6] == ["Adam Durand", "Hugo Petit", "Léa Bernard", "P0 Nom00", "P1 Nom01", "P2 Nom02"],
          f"list view: alphabetical as displayed (first name, then last name; P2 before P10) {names()[:6]}")
    page.evaluate("document.getElementById('tab-contacts').scrollTop=600")
    before = page.evaluate("document.getElementById('tab-contacts').scrollTop")
    row = page.locator("#contacts-list .contact-row", has_text="Nom20")
    row.locator(".contact-main").click()
    page.locator(".contact-detail-links button", has_text="Modifier la fiche").click()
    check(page.is_visible("#contact-form") and page.evaluate("document.getElementById('tab-contacts').scrollTop") == before, "opening the edit form keeps the list position")
    page.fill("#ct-role", "Acheteur"); page.click("#contact-save-btn")
    check(page.evaluate("document.getElementById('tab-contacts').scrollTop") == before, "saving a contact keeps the list position")
    check(page.evaluate("document.querySelector('.app-sidebar').getBoundingClientRect().top") == 0, "sidebar stays in place")
    page.evaluate("window.scrollTo(0, 2000)")
    check(page.evaluate("document.querySelector('.app-sidebar').getBoundingClientRect().top") == 0, "sidebar stays in place when the page scrolls")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline: client → contact → mail, installations grouped in a pricing request, Go, Excel rows
    MAIL = ("De : Pierre Lacombe <p.lacombe@vergersdulot.fr>\nObjet : Projet centrale PV\n\n"
            "Nous mettrons en service en juin 2027 une centrale photovoltaïque au sol de 2,4 MWc à Prayssac (46220), "
            "lauréate de l'appel d'offres simplifié de la CRE. Production attendue : environ 3 100 MWh par an.\n\n"
            "Cordialement,\nPierre Lacombe\nGérant\nSCEA LES VERGERS DU LOT\n06 71 23 45 89")
    legacy = [{"id": "o1", "company": "Solaris", "title": "AO Q3", "stage": "pricing", "taskId": None, "createdAt": 1, "updatedAt": 1, "deleted": False}]
    anne = [{"id": "c1", "firstName": "Anne", "lastName": "Roux", "role": "Asset manager", "company": "Solaris", "email": "anne@urba.fr",
             "phone1": "0611223344", "phone2": "", "relevance": 3, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(legacy), "todo_todo-widget-contacts": json.dumps(anne),
                                      "todo_todo-widget-clients": json.dumps(["Solaris"])}, viewport={"width": 1300, "height": 850})
    page.context.grant_permissions(["clipboard-read", "clipboard-write"])
    opps = lambda: stored(page, "todo-widget-opportunities")
    page.goto(URL); page.click('[data-tab="pipeline"]')
    check("Demande de prix" in page.inner_text('[data-phase="offre"]'), "old opportunities move to the new stages")
    page.click("#opp-add-btn"); page.fill("#pq1-company", "solaris")
    check("Client existant" in page.inner_text("#pq1-company-hint"), "step 1 recognises a known client")
    page.click("#pq-next-btn")
    check(page.input_value("#pq-s-mail") == "anne@urba.fr", "its only contact is used by default")
    page.click("#pq-cancel-btn")
    page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA Les Vergers du Lot")
    check("Nouveau client" in page.inner_text("#pq1-company-hint"), "step 1 flags a new client")
    page.click("#pq-next-btn"); page.fill("#pq-mail", MAIL); page.wait_for_timeout(500)
    check(page.input_value("#pq-0-kwc") == "2400" and page.input_value("#pq-0-dept") == "46 - Lot" and page.input_value("#pq-s-prenom") == "Pierre",
          "pasted mail prefills the installation and the contact")
    page.click("#pq-add-inst"); page.fill("#pq-1-centrale", "Le Moulin"); page.fill("#pq-1-kwc", "1800")
    page.select_option("#pq-s-charge", "PGU"); page.click("#pq-save-btn")
    check(page.is_visible("#pq-overlay") and "nombre d" in page.inner_text("#pq-summary").lower(), "two installations: saving waits for the count to be confirmed")
    check("2 installations" in page.inner_text("#pq-count-box"), "the form shows the number of installations found")
    page.check("#pq-count-ok"); page.click("#pq-save-btn")
    inst = [o for o in opps() if o.get("kind") == "installation" and o["company"] == "SCEA Les Vergers du Lot"]
    req = [o for o in opps() if o.get("kind") == "request"]
    check(len(inst) == 2 and len(req) == 1 and all(o["requestId"] == req[0]["id"] for o in inst), "two installations grouped in one request")
    check(any(c["email"] == "p.lacombe@vergersdulot.fr" for c in stored(page, "todo-widget-contacts")), "contact created from the mail")
    t = [t for t in stored(page, "todo-widget-tasks") if t.get("opportunityId") == req[0]["id"]]
    check(len(t) == 1 and t[0]["title"].startswith("Appeler Pierre"), "one call task for the whole request")
    page.click(".pipe-req-card .pipe-card-title")
    check("Codes affaire" in page.input_value("#opp-mail-body") and "27(06)" in page.input_value("#opp-mail-body"), "pricing mail with business codes")
    page.click("#opp-detail-actions >> text=Copier lignes Pipe")
    page.click("#choice-btns >> text=Une ligne par installation")
    rows = page.evaluate("navigator.clipboard.readText()").split("\n")
    check(len(rows) == 2 and all(len(r.split("\t")) == 37 for r in rows), "Pipe rows: one per installation, 37 columns")
    page.click("#opp-detail-actions >> text=Fermer")
    # validating the task offers the next stage for the whole request
    page.click('[data-tab="tasks"]'); page.locator(".task-line", has_text="Appeler Pierre").locator(".task-check").first.click()
    page.wait_for_timeout(200)
    check(page.is_visible("#stage-advance-overlay"), "validating a linked task offers the next stage")
    page.click("#stage-advance-confirm")
    check(all(o["stage"] == "prise_contact" for o in opps() if o.get("requestId") == req[0]["id"]), "the request's installations advance together")
    # Go from the installation sheet
    page.click('[data-tab="pipeline"]')
    page.locator('[data-phase="qualif"] .pipe-req-name', has_text="Le Moulin").click()
    page.click(".opp-stage-step:has-text('Qualification')")
    page.click("#opp-detail-actions >> text=Go"); page.click("#stage-advance-confirm")
    moulin = next(o for o in opps() if o.get("centrale") == "Le Moulin")
    check(moulin["stage"] == "demande_prix" and moulin.get("goDate"), "Go moves to the pricing request and records the date")
    check(page.is_visible("#opp-mail-body"), "after Go, the pricing request opens")
    page.click("#opp-detail-actions >> text=Fermer")
    page.locator('[data-phase="offre"] .pipe-card', has_text="Le Moulin").first.click()
    check("statut bdd" in page.inner_text("#opp-detail-body").lower() and "En préparation" in page.inner_text("#opp-detail-body"), "installation sheet shows the BDD status")
    page.click("#opp-detail-actions >> text=Copier ligne BDD")
    row = page.evaluate("navigator.clipboard.readText()").split("\t")
    check(len(row) == 91 and row[24] == "En préparation" and row[33] == "Le Moulin", "BDD row: 91 columns, status and plant")
    check(re.match(r"^\[?\w*\]?[\w()\[\]]* – .+_Le Moulin – A$", row[7]) is not None, f"BDD row: business code in column H « Ref_Watt.ee » ({row[7]})")
    page.click("#opp-detail-actions >> text=Fermer")
    # pré-qualif → installation
    page.click("#prequal-add-btn"); page.fill("#prequal-company", "Ferme Dupont"); page.fill("#prequal-note", "Salon"); page.click("#prequal-save-btn")
    page.click('[data-phase="prequal"] .pipe-card'); page.click("#opp-detail-actions >> text=Projet identifié")
    check(page.input_value("#pq1-company") == "Ferme Dupont", "pré-qualif opens the form on its client")
    page.click("#pq-next-btn"); page.fill("#pq-0-centrale", "Toiture A"); page.click("#pq-save-btn")
    check(page.locator('[data-phase="prequal"] .pipe-card').count() == 0 and "Toiture A" in page.inner_text('[data-phase="qualif"]'), "pré-qualif becomes an installation")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Phone version (index.html?lite): bottom bar, fewer features, pipeline read-only
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps([{"id": "i3", "kind": "installation", "company": "Solaris",
        "centrale": "Moulin Neuf", "title": "Moulin Neuf", "stage": "negociation", "createdAt": 1, "updatedAt": 1, "deleted": False}])},
        viewport={"width": 390, "height": 800}, is_mobile=True, has_touch=True)
    page.goto(URL + "?lite")
    check(page.is_visible("#lite-nav") and not page.is_visible(".app-sidebar"), "phone version: bottom bar instead of the sidebar")
    check(page.get_attribute("#app-manifest", "href") == "manifest-lite.json", "phone version installs with its own manifest")
    page.click("#lite-nav [data-go=tasks]"); page.fill("#quick-add-input", "Depuis le téléphone"); page.press("#quick-add-input", "Enter")
    check(any(t["title"] == "Depuis le téléphone" for t in stored(page, "todo-widget-tasks")), "phone version: quick add works")
    page.click("#lite-nav [data-go=pipeline]")
    check(not page.is_visible("#opp-add-btn"), "phone version: no opportunity form")
    page.click(".pipe-card:has-text('Moulin Neuf')")
    visible = page.eval_on_selector_all("#opp-detail-actions button", "els=>els.filter(e=>e.offsetParent).map(e=>e.textContent)")
    check(visible == ["Fermer"], f"phone version: installation sheet is read-only ({visible})")
    check(page.evaluate("document.documentElement.scrollWidth") <= 390, "phone version fits the screen")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Contacts: « Coller une signature » pre-fills the contact form
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="contacts"]'); page.click("#sig-paste-btn")
    page.fill("#sig-paste-text", "M. Jean DUPONT\nPrésident\nSAS SOLAIRE DU MIDI\nTél : 04 67 11 22 33 - Mob. : 06 12 34 56 78\njean.dupont@solairedumidi.fr")
    page.click("#sig-paste-analyze")
    got = [page.input_value(s_) for s_ in ("#ct-first", "#ct-last", "#ct-role", "#ct-company", "#ct-email", "#ct-phone1", "#ct-phone2")]
    check(got == ["Jean", "Dupont", "Président", "SAS SOLAIRE DU MIDI", "jean.dupont@solairedumidi.fr", "06 12 34 56 78", "04 67 11 22 33"], f"signature fills the contact form ({got})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- New opportunity: pasting a signature in step 1 fills the client and the contact
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn")
    page.click("#pq1-sig-box summary")
    page.fill("#pq1-sig", "Camille DUBOIS\nResponsable financement de projets\nRHÔNE ÉNERGIES SAS\nT. 04 75 00 00 06\nM. 06 77 00 00 06\ncamille.dubois@rhone-energies.fr")
    page.wait_for_timeout(500)
    got = [page.input_value(x) for x in ("#pq1-company", "#pq1-prenom", "#pq1-nom", "#pq1-qualite", "#pq1-tel", "#pq1-tel2", "#pq1-mail")]
    check(got == ["RHÔNE ÉNERGIES SAS", "Camille", "DUBOIS", "Responsable financement de projets", "06 77 00 00 06", "04 75 00 00 06", "camille.dubois@rhone-energies.fr"],
          f"step 1: a pasted signature fills client and contact ({got})")
    # a second signature (another person) becomes an additional contact; the main one stays
    page.fill("#pq1-sig", "Marc LEROY\nDirecteur technique\nRHÔNE ÉNERGIES SAS\n06 11 11 11 11\nmarc.leroy@rhone-energies.fr"); page.wait_for_timeout(500)
    check(page.input_value("#pq1-prenom") == "Camille" and page.input_value("#pq1-extras input[data-xk='prenom']") == "Marc", "a second signature adds a second contact")
    page.click("#pq1-add-extra"); page.fill("#pq1-extras input[data-x='1'][data-xk='prenom']", "Lucie"); page.fill("#pq1-extras input[data-x='1'][data-xk='nom']", "Roy")
    check(page.locator("#pq1-principal [data-principal]").count() == 3, "the main contact can be chosen among the three people")
    page.click("#pq1-principal [data-principal='x:0']")   # Marc becomes the main contact
    page.click("#pq-next-btn")
    check(page.input_value("#pq-s-producteur") == "RHÔNE ÉNERGIES SAS" and page.input_value("#pq-s-prenom") == "Marc" and page.input_value("#pq-s-tel") == "06 11 11 11 11",
          "step 2 uses the chosen main contact")
    page.select_option("#pq-contact-switch", "p:0"); page.wait_for_timeout(100)
    check(page.input_value("#pq-s-prenom") == "Camille" and page.input_value("#pq-s-tel") == "06 77 00 00 06 / 04 75 00 00 06", "step 2: the main contact can be changed")
    page.select_option("#pq-contact-switch", "p:1"); page.wait_for_timeout(100)
    page.fill("#pq-0-centrale", "Toit A"); page.click("#pq-save-btn")
    names = sorted(c["firstName"] for c in stored(page, "todo-widget-contacts"))
    inst = [o for o in stored(page, "todo-widget-opportunities") if o.get("kind") == "installation"][0]
    check(names == ["Camille", "Lucie", "Marc"] and len(inst.get("otherContactIds", [])) == 2, f"three contacts created, two linked as other contacts ({names})")
    camille = next(c for c in stored(page, "todo-widget-contacts") if c["firstName"] == "Camille")
    check(camille["phone1"] == "06 77 00 00 06" and camille["phone2"] == "04 75 00 00 06" and camille["id"] in inst.get("otherContactIds", []),
          "the first contact, no longer main, keeps both numbers and is linked as an other contact")
    check(inst.get("prenom") == "Marc", "the installation's main contact is the chosen one")
    # editing: switch the main contact to Camille, complete an empty field from a pasted mail (filled fields untouched)
    page.locator(".pipe-card", has_text="Toit A").first.click(); page.click("#opp-detail-edit")
    cam_opt = page.eval_on_selector("#pq-contact-switch", "s=>[...s.options].find(o=>o.textContent.startsWith('Camille')).value")
    page.select_option("#pq-contact-switch", cam_opt); page.wait_for_timeout(100)
    page.fill("#pq-mail", "Bonjour, notre centrale fait 2,4 MWc à Valence (26000).\nCordialement"); page.wait_for_timeout(500)
    check(page.input_value("#pq-0-kwc") == "2400" and page.input_value("#pq-0-centrale") == "Toit A", "editing: a pasted mail fills empty fields only")
    page.click("#pq-save-btn")
    inst = [o for o in stored(page, "todo-widget-opportunities") if o.get("kind") == "installation"][0]
    marc = next(c for c in stored(page, "todo-widget-contacts") if c["firstName"] == "Marc")
    check(inst.get("prenom") == "Camille" and marc["id"] in inst.get("otherContactIds", []) and camille["id"] not in inst.get("otherContactIds", []),
          "editing: the main contact changes, the previous one stays as an other contact")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Qualification: the number of plants announced by the mail is checked against the installations found
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "SEM"); page.click("#pq-next-btn")
    page.fill("#pq-mail", "De : H R <h@sem.fr>\nObjet : Nos centrales\n\nNous avons trois projets :\n- Gymnase : 250 kWc\n- École : 100 kWc\nSoit une puissance totale de 600 kWc.\nH R")
    page.wait_for_timeout(500)
    box = page.inner_text("#pq-count-box")
    check("annonce 3 installations" in box and "600 kWc" in box, f"missing plant flagged against the mail ({box!r})")
    page.click("#pq-add-inst"); page.fill("#pq-2-centrale", "Mairie"); page.fill("#pq-2-kwc", "250"); page.press("#pq-2-kwc", "Tab"); page.wait_for_timeout(200)
    check("conforme au mail" in page.inner_text("#pq-count-box"), "adding the missing plant makes the count match")
    page.check("#pq-count-ok"); page.click("#pq-save-btn")
    check(len([o for o in stored(page, "todo-widget-opportunities") if o.get("kind") == "installation"]) == 3, "three installations saved after confirmation")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline: a request whose installations are at different stages moves each one forward (never back)
    mixed = [{"id": "r9", "kind": "request", "company": "SEM", "title": "Demande de pricing", "createdAt": 1, "updatedAt": 1, "deleted": False},
             {"id": "a", "kind": "installation", "company": "SEM", "centrale": "A", "title": "A", "stage": "negociation", "requestId": "r9", "createdAt": 1, "updatedAt": 1, "deleted": False},
             {"id": "b", "kind": "installation", "company": "SEM", "centrale": "B", "title": "B", "stage": "collecte", "requestId": "r9", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(mixed)})
    page.goto(URL); page.click('[data-tab="pipeline"]')
    page.locator('[data-phase="contrat"] .pipe-card', has_text="A").first.click()
    page.click("#opp-detail-body .opp-back")
    page.click("#opp-detail-actions >> text=Avancer chaque installation"); page.click("#stage-advance-confirm")
    st = {o["id"]: o.get("stage") for o in stored(page, "todo-widget-opportunities")}
    check(st["a"] == "pricing_ferme" and st["b"] == "qualif_en_cours", f"mixed request: each installation to its own next stage ({st})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline: offer sent → waiting for the producer's answer; accepted (celebration) or refused (lost, its own animation)
    wait = [{"id": "r8", "kind": "request", "company": "SEM", "title": "Demande de pricing", "createdAt": 1, "updatedAt": 1, "deleted": False}] + [
            {"id": k, "kind": "installation", "company": "SEM", "centrale": k.upper(), "title": k.upper(), "stage": "offre_remise", "requestId": "r8",
             "createdAt": 1, "updatedAt": 1, "deleted": False} for k in ("w1", "w2")] + [
            {"id": "solo", "kind": "installation", "company": "ZEN", "centrale": "SOLO", "title": "SOLO", "stage": "offre_attente", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(wait)})
    page.goto(URL); page.click('[data-tab="pipeline"]')
    page.locator(".pipe-req-name", has_text="W1").click(); page.click("#opp-detail-body .opp-back")
    page.click("#opp-detail-actions >> text=Tout passer à l\u2019étape suivante")
    check("Retour sur l\u2019offre" in page.inner_text("#stage-advance-text"), "after « offre à remettre » comes the wait for the producer's answer")
    page.click("#stage-advance-confirm")
    st = {o["id"]: o.get("stage") for o in stored(page, "todo-widget-opportunities")}
    check(st["w1"] == st["w2"] == "offre_attente", f"both installations wait for the answer ({st})")
    t = [t for t in stored(page, "todo-widget-tasks") if t.get("opportunityId") == "r8"]
    check(len(t) == 1 and t[0]["status"] == "en_attente" and "Relancer" in t[0]["title"], "a follow-up task « en attente » to chase the producer")
    page.locator(".pipe-req-card", has_text="W1").locator(".pipe-card-title").first.click()
    page.click("#opp-detail-actions >> text=Offre acceptée (2)"); page.click("#stage-advance-confirm")
    page.wait_for_selector(".celebrate-won", timeout=2000)
    check(page.locator(".celebrate-won .confetti").count() > 20 and "Offre acceptée" in page.inner_text(".celebrate-won") and page.locator(".celebrate").count() == 1,
          "offer accepted: one confetti animation for the whole request")
    check(all(o["stage"] == "kyc" for o in stored(page, "todo-widget-opportunities") if o.get("requestId") == "r8"), "accepted: the installations move to KYC")
    page.locator('[data-phase="contrat"] .pipe-card').first.click()   # the animation never blocks a click
    page.click("#opp-detail-actions >> text=Fermer")
    page.wait_for_selector(".celebrate", state="detached", timeout=5000)
    page.once("dialog", lambda d: d.accept())
    page.locator(".pipe-card", has_text="SOLO").first.click(); page.click("#opp-detail-actions >> text=Offre refusée")
    page.wait_for_selector(".celebrate-lost", timeout=2000)
    check("Offre perdue" in page.inner_text(".celebrate-lost") and next(o for o in stored(page, "todo-widget-opportunities") if o["id"] == "solo")["stage"] == "perdu",
          "offer refused: lost, with its own animation")
    seq = page.evaluate("""(()=>{ const out={}; ['won','signed','lost'].forEach(k=>{ out[k]=[]; for(let i=0;i<30;i++){ window.todoCelebrate(k, {sub:'x'}); out[k].push(document.querySelector('.celebrate').dataset.variant); } }); return out; })()""")
    check(all(len(set(v)) == 3 and all(a_ != b_ for a_, b_ in zip(v, v[1:])) for v in seq.values()),
          f"animations: three per moment, drawn at random, never the same twice in a row ({seq['signed'][:6]})")
    shown = page.evaluate("""(()=>{ const r=[]; [['won','panel'],['won','turbine'],['won','battery'],['signed','gun'],['signed','gong'],['signed','champagne'],['lost','umbrella'],['lost','nowind'],['lost','plane']].forEach(([k,v])=>{
      window.todoCelebrate(k, {sub:'x'}, v); const el=document.querySelector('.celebrate'); r.push(el.dataset.variant===v && el.getBoundingClientRect().width===innerWidth && !!el.querySelector('.cx-msg b').textContent); }); return r; })()""")
    check(all(shown), f"animations: the nine play full screen with their message ({shown})")
    page.evaluate("document.querySelectorAll('.celebrate').forEach(n=>n.remove())")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline · meeting view: one line per installation by phase, copied for a mail, as an Excel file
    def mi(i, co, n, st, kwc, mwh, mes, tech="PV"):
        return {"id": "mt%d" % i, "kind": "installation", "company": co, "producteur": co, "centrale": n, "title": n, "stage": st, "kwc": kwc, "mwh": mwh, "mesPrev": mes,
                "install": tech, "contrat": "AO CRE CR", "fin": "31/12/2030", "dept": "49 - Maine-et-Loire", "createdAt": 1, "updatedAt": 1, "deleted": False}
    meet = [mi(1, "SOLARIS", "Toit Durand", "collecte", "300", "360", "01/06/2027"), mi(2, "SOLARIS", "CS du Moulin", "demande_prix", "4500", "5200", "01/03/2027"),
            mi(3, "VENTUS", "Parc des Crêtes", "offre_attente", "12000", "28000", "01/12/2027", "Eolien"), mi(4, "ZEN", "Hangar 2", "negociation", "250", "290", "15/09/2026"),
            dict(mi(5, "ZEN", "Perdu", "perdu", "100", "", "2027"))]
    ctx = browser.new_context(viewport={"width": 1400, "height": 900}, permissions=["clipboard-read", "clipboard-write"], accept_downloads=True)
    ctx.add_init_script("if(!sessionStorage.getItem('seeded')){ localStorage.setItem('todo_todo-widget-opportunities', %s); sessionStorage.setItem('seeded','1'); }" % json.dumps(json.dumps(meet)))
    page = ctx.new_page(); errors = []; page.on("pageerror", lambda e: errors.append(str(e)))
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#pipe-meet-toggle")
    secs = page.eval_on_selector_all(".meet-table tr.meet-sec th", "els=>els.map(e=>e.innerText.replace(/\\s+/g,' ').trim())")
    names = page.eval_on_selector_all(".meet-table td.meet-name", "els=>els.map(e=>e.textContent)")
    check(len(secs) == 3 and names == ["Toit Durand", "Parc des Crêtes", "CS du Moulin", "Hangar 2"] and "4,5 MWc" in page.inner_text(".meet-table") and "jusqu’au 31/12/2030" in page.inner_text(".meet-table"),
          f"meeting view: one line per installation, by phase, the most advanced first, lost ones left out ({secs}, {names})")
    page.click(".meet-chip[data-phase='qualif']")
    check(page.locator(".meet-table td.meet-name").count() == 3 and page.get_attribute(".meet-chip[data-phase='qualif']", "aria-pressed") == "false", "meeting view: a phase can be left out")
    page.click(".meet-actions >> text=Copier pour mail / Teams"); page.wait_for_timeout(200)
    clip = page.evaluate("navigator.clipboard.readText()")
    check(clip.startswith("Pipe agrégation · ce qui arrive") and "Parc des Crêtes" in clip and "Toit Durand" not in clip, "meeting view: copied for a mail (the phases shown only)")
    with page.expect_download() as dl:
        page.click(".meet-actions >> text=Excel")
    f = pathlib.Path(tempfile.mkdtemp()) / dl.value.suggested_filename; dl.value.save_as(f)
    ws = openpyxl.load_workbook(f).active
    vals = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    check(f.name.endswith("_Pipe agrégation - réunion pricing.xlsx") and vals[2][3] == "Puissance (kWc)" and any(r[1] == "CS du Moulin" and r[3] == 4500 for r in vals),
          f"meeting view: Excel file with real numbers ({f.name})")
    page.click("#pipe-meet-toggle"); check(page.locator(".pipe-col").count() > 0, "meeting view: back to the board")
    check(not errors, f"no page errors ({errors})")
    ctx.close()

    # --- Pipeline: a mail listing several plants gives one installation each; quotes in a mail never reach the HTML
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "HELIO OUEST SAS"); page.click("#pq-next-btn")
    page.fill("#pq-mail", "De : R C <r@helio.fr>\nObjet : x\n\nBonjour,\n- Centrale de Cholet (49) : 1,2 MWc, MES 03/2027\n- Centrale de Saumur (49) : 800 kWc\nCordialement,\nRomain Chevalier")
    page.wait_for_timeout(500)
    check(page.input_value("#pq-1-centrale") == "Centrale de Saumur" and page.input_value("#pq-0-mesPrev") == "Mars 2027", "mail listing two plants → two installations")
    page.fill("#pq-mail", "De\tObjet\tReçu\tTaille\tCatégories\t\nPierre Dupont;  Paul Guérin\tRE: Projets PV\tlun. 16:07\t318 Ko\tTodo-synced"); page.wait_for_timeout(400)
    check(page.is_visible("#pq-mail-warning"), "pasting the Outlook list line explains how to get the mail content")
    page.fill("#pq-0-centrale", 'x" autofocus onfocus="window.__xss=1" y="')
    page.select_option("#pq-0-install", "Eolien"); page.wait_for_timeout(300)  # re-renders the form with the value in an attribute
    check(page.evaluate("window.__xss") is None and page.input_value("#pq-0-centrale").startswith('x" autofocus'), "a quote in a field can't inject HTML")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Pipeline: producer Excel, one column per plant (SheetJS from cdnjs)
    wb = openpyxl.Workbook(); ws = wb.active
    for r in [["", "Centrale 1", "Centrale 2"], ["Nom Installation", "Les Vignes", "Le Moulin"], ["Département", "Lot (46)", "Gers (32)"],
              ["Puissance installée (MWc)", 2.4, 1.85], ["P50 (MWh/an)", 3100, 2300], ["Date de mise en service", "Juin 2027", "Mars 2028"]]:
        ws.append(r)
    xlsx = pathlib.Path(tempfile.mkdtemp()) / "carac.xlsx"; wb.save(xlsx)
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA X"); page.click("#pq-next-btn")
    page.set_input_files("#pq-file", str(xlsx)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    check(page.input_value("#pq-0-centrale") == "Les Vignes" and page.input_value("#pq-1-kwc") == "1850" and page.input_value("#pq-1-dept") == "32 - Gers",
          "producer Excel creates one installation per column")
    # the Selfee « Caractéristiques Centrales » template: SPV, SIRET, P90 reach the form, kVA/productible rows are not taken as power
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Données"
    for r in [["", "Centrale 1"], ["Nom Installation", "CS Bram"], ["Nom SPV", "BRAM SOLAIRE SAS"], ["SIRET SPV", "123 456 789 00012"], ["Département", "Aude (11)"],
              ["Puissance raccordement (kVA)", 4200], ["Puissance installée (kWc)", 4986.93], ["P50 (MWh/an)", 6650], ["P90 (MWh/an)", 6100], ["Productible P50 (kWh/kWc/an)", 1333]]:
        ws.append(r)
    tpl = pathlib.Path(tempfile.mkdtemp()) / "carac_selfee.xlsx"; wb.save(tpl)
    page.click("#pq-cancel-btn"); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA Y"); page.click("#pq-next-btn")
    page.set_input_files("#pq-file", str(tpl)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    got = [page.input_value(s_) for s_ in ("#pq-0-centrale", "#pq-0-kwc", "#pq-0-spv", "#pq-0-siret", "#pq-0-p90", "#pq-0-dept")]
    check(got == ["CS Bram", "4986.93", "BRAM SOLAIRE SAS", "123 456 789 00012", "6100", "11 - Aude"], f"Selfee template read into the form ({got})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- A mail thread with several people: all detected, ticked or not, any of them can become the main contact
    thread = """De : Jean Martin <j.martin@exemple-solaire.fr>
Envoyé : mardi 20 octobre 2026 10:12
À : Paul Collègue <paul.collegue@selfee.fr>
Cc : Sophie Durand <s.durand@exemple-solaire.fr>
Objet : RE: Centrale de Villers

Bonjour,
Voici les éléments pour la centrale de Villers, 450 kWc.
Cordialement,
Jean MARTIN
Gérant
06 10 20 30 40

De : Paul Collègue <paul.collegue@selfee.fr>
Envoyé : lundi 19 octobre 2026 09:00
À : Paul Bernard <p.bernard@exemple-solaire.fr>
Objet : Centrale de Villers

Bonjour, pouvez-vous m'envoyer les caractéristiques ?

De : Paul Bernard <p.bernard@exemple-solaire.fr>
Envoyé : vendredi 16 octobre 2026 17:40
À : contact@selfee.fr
Objet : Centrale de Villers

Bonjour, nous cherchons un agrégateur pour notre centrale.
Bien à vous,
Paul BERNARD
Responsable technique
06 22 33 44 55
"""
    page, errors = new_page(browser, viewport={"width": 1300, "height": 900})
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "EXEMPLE SOLAIRE"); page.click("#pq-next-btn")
    page.fill("#pq-mail", thread); page.wait_for_selector("#pq-detect:not([hidden]) .pq-person", timeout=5000)
    people = page.eval_on_selector_all("#pq-detect .pq-person", "els=>els.map(e=>[e.querySelector('b').textContent, e.querySelector('input').checked, e.querySelector('em').textContent])")
    names = sorted(p[0] for p in people)
    check(names == ["Jean MARTIN", "Paul BERNARD", "Sophie DURAND"] and dict((p[0], p[1]) for p in people)["Sophie DURAND"] is False
          and all(p[1] for p in people if p[0] != "Sophie DURAND"), f"import: every person of the thread found, colleagues left out, people in copy not ticked ({people})")
    check("1 installation" in page.inner_text("#pq-detect") and "450 kWc" in page.inner_text("#pq-detect"), "import: what was found is summed up (plants)")
    page.locator("#pq-detect .pq-person", has_text="Sophie DURAND").locator(".pq-star").click()
    check(page.input_value("#pq-s-prenom") == "Sophie" and page.input_value("#pq-s-nom") == "DURAND", "import: ★ makes a detected person the main contact")
    page.fill("#pq-0-centrale", "Villers")
    page.click("#pq-save-btn"); page.wait_for_timeout(300)
    if page.is_visible("#pq-count-ok"): page.check("#pq-count-ok"); page.click("#pq-save-btn"); page.wait_for_timeout(300)
    cts_now = {c["firstName"] + " " + c["lastName"]: c["id"] for c in stored(page, "todo-widget-contacts") if not c.get("deleted")}
    opp = next(o for o in stored(page, "todo-widget-opportunities") if o.get("kind") == "installation")
    linked = sorted(k for k, v in cts_now.items() if v in (opp.get("otherContactIds") or []))
    check(sorted(cts_now) == ["Jean Martin", "Paul Bernard", "Sophie Durand"] and opp.get("nom") == "DURAND" and linked == ["Jean Martin", "Paul Bernard"],
          f"import: the ticked people are created and linked, the main one in the Pipe columns ({sorted(cts_now)}, linked {linked})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Editing one installation with a producer workbook of several plants (no plant names, SPV only, summer dates)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Données"
    for r in [["", "Centrale 1", "Centrale 2", "Centrale 3"], ["Nom Installation"], ["Nom SPV", "ALPHA PV1", "BETA PV1"], ["SIRET SPV", 12345678900011, 98765432100022],
              ["Département", 49, 49], ["Commune", "49640 VILLE-SUR-EXEMPLE", "49330 AUTRE-VILLE"], ["Nature Projet", "Centrale au sol", "Toiture"],
              ["Date de mise en service (effective ou prévisionnelle)", datetime.datetime(2027, 7, 20), datetime.datetime(2027, 7, 20)],
              ["Puissance installée (kWc)", "299", "211,14"], ["P50 annuel (indicatif) (MWh/an)", 358, 252]]:
        ws.append(r)
    multi = pathlib.Path(tempfile.mkdtemp()) / "caracteristiques.xlsx"; wb.save(multi)
    ed_inst = [{"id": "e1", "kind": "installation", "company": "SOLARIS", "producteur": "SOLARIS", "centrale": "", "title": "", "kwc": "299", "stage": "collecte",
                "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(ed_inst)}, viewport={"width": 1300, "height": 900}, timezone_id="Europe/Paris")
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.locator(".pipe-card").first.click(); page.click("#opp-detail-edit")
    page.set_input_files("#pq-file", str(multi)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    got = [page.input_value(s_) for s_ in ("#pq-0-centrale", "#pq-0-spv", "#pq-0-cp", "#pq-0-commune", "#pq-0-mesPrev")]
    check(got == ["ALPHA PV1", "ALPHA PV1", "49640", "VILLE-SUR-EXEMPLE", "20/07/2027"],
          f"edit + workbook: the matching plant (299 kWc) fills the form, SPV names it, postal code split, summer date right in Paris time ({got})")
    heads = page.eval_on_selector_all("#pq-form .pq-inst .pq-sec-h", "els=>els.map(e=>e.innerText.replace(/\\s+/g,' ').trim())")
    check(page.input_value("#pq-1-centrale") == "BETA PV1" and page.input_value("#pq-1-kwc") == "211.14" and page.input_value("#pq-1-mesPrev") == "20/07/2027"
          and "nouvelle : sera créée" in heads[1] and "celle que tu modifies" in heads[0],
          f"edit + workbook: the other plant shown right away in the form as Installation 2, to check ({heads})")
    page.click("#pq-save-btn"); page.wait_for_timeout(300)
    opps_now = [o for o in stored(page, "todo-widget-opportunities") if not o.get("deleted")]
    insts = sorted((o for o in opps_now if o["kind"] == "installation"), key=lambda o: o["kwc"])
    reqs = [o for o in opps_now if o["kind"] == "request"]
    check(len(insts) == 2 and len(reqs) == 1 and all(o["requestId"] == reqs[0]["id"] for o in insts) and insts[0]["centrale"] == "BETA PV1"
          and insts[0]["kwc"] == "211.14" and insts[0]["stage"] == "collecte", f"edit + workbook: the second plant created in the same request ({[o.get('centrale') for o in insts]})")
    check(insts[1]["mesPrev"] == "20/07/2027" and insts[1]["mwh"] == "358", "edit + workbook: commissioning date and volume saved in the installation")
    # editing again with the same workbook: both installations shown, matched, nothing created twice
    page.wait_for_selector("#opp-detail-overlay:not(.hidden)"); page.click("#opp-detail-edit")   # the sheet reopens after saving
    check(page.input_value("#pq-1-centrale") == "BETA PV1" and "déjà dans la demande" in page.inner_text("#pq-form"), "edit again: the request's other installation is in the form")
    page.set_input_files("#pq-file", str(multi)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    check(page.locator("#pq-form .pq-inst").count() == 2, "edit again: the workbook's plants match the existing installations (no third one)")
    page.click("#pq-save-btn"); page.wait_for_timeout(300)
    n_inst = len([o for o in stored(page, "todo-widget-opportunities") if o["kind"] == "installation" and not o.get("deleted")])
    check(n_inst == 2, f"edit again: still two installations ({n_inst})")
    page, errors = new_page(browser)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA Y"); page.click("#pq-next-btn")
    # a mail saved as .eml (new Outlook, Outlook web): its text and its attached Excel, also inside a forwarded mail
    from email.message import EmailMessage
    inner = EmailMessage(); inner["From"] = "Jean Martin <j.martin@exemple-solaire.fr>"; inner["Subject"] = "Nos centrales"
    inner.set_content("Bonjour,\nVoici la liste de nos centrales.\nJean MARTIN\nGérant\n06 10 20 30 40")
    inner.add_attachment(xlsx.read_bytes(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="Caractéristiques centrales.xlsx")
    outer = EmailMessage(); outer["From"] = "Paul Collègue <paul@exemple-agregateur.fr>"; outer["Subject"] = "TR: Nos centrales"
    outer.set_content("Pour info, à qualifier.")
    outer.add_attachment(inner)
    eml = pathlib.Path(tempfile.mkdtemp()) / "mail.eml"; eml.write_bytes(bytes(outer))
    page.click("#pq-cancel-btn"); page.click("#opp-add-btn"); page.fill("#pq1-company", "SCEA Z"); page.click("#pq-next-btn")
    page.set_input_files("#pq-file", str(eml)); page.wait_for_function("document.getElementById('pq-files-info').textContent.includes('installation')", timeout=20000)
    check(page.input_value("#pq-0-centrale") == "Les Vignes" and page.input_value("#pq-1-kwc") == "1850" and "Caractéristiques centrales.xlsx" in page.inner_text("#pq-files-info"),
          "a .eml's attached Excel is read, even in a forwarded mail (accented file name kept)")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Parc: Suivi Agrégation workbook → dashboard; contacts → phone (.vcf) without duplicates
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BDD Agrégation"
    hdr = ["Vérification", "Année", "Chargé", "Date entrée BDD", "Source", "Crédit Agricole", "Développeur", "Ref", "Cadre juridique", "Type AO CRE",
           "Région", "Département", "Statut", "Date_fin", "Renouvellement", "Nom_prod", "Nom_Installation_prod", "Type_EnR", "P_installée (KWc)", "P50\n(MWh)",
           "Agrégation 2026", "Agrégation 2027", "Département 2", "Département2"]
    ws.append(["Date renouvellement"]); ws.append([]); ws.append(hdr)
    soon = (TODAY + datetime.timedelta(days=90))
    ws.append(["", 2021, "PGU", None, "Contact direct", "N/A", "SOLARIS", "", "CR", "AO CRE", "Occitanie", "34 - Hérault", "Gagné", soon, "A renouveler", "SOLARIS", "CS DU MOULIN", "PV", 2990, 4500, 4500, 0, "Hérault", "30 - Gard"])
    ws.append(["", 2022, "SMA", None, "Contact direct", "N/A", "VENTUS", "", "Marché", "", "Bretagne", "29 - Finistère", "Gagné", datetime.date(2028, 12, 31), "En cours", "VENTUS", "Parc B", "PV", 10000, 23000, 23000, 23000])
    ws.append(["", 2026, "PGU", None, "BOAMP", "N/A", "X", "", "CR", "", "Grand Est", "", "En attente", None, "", "X", "Toit", "PV", 500, 600, "", ""])
    d = wb.create_sheet("Dashboard"); d.append(["", "Désignation", "2026", "2027"]); d.append(["", "Object. Agrégation Vol.", 150, 300])
    parc = pathlib.Path(tempfile.mkdtemp()) / "suivi.xlsx"; wb.save(parc)
    dupes = [{"id": "c1", "firstName": "Anne", "lastName": "Roux", "role": "Asset", "company": "Solaris", "email": "anne@urba.fr", "phone1": "0611223344", "phone2": "",
              "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 1, "deleted": False},
             {"id": "c2", "firstName": "Anne", "lastName": "Roux", "role": "", "company": "Solaris", "email": "ANNE@urba.fr", "phone1": "", "phone2": "",
              "relevance": 0, "comment": "", "createdAt": 1, "updatedAt": 2, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-contacts": json.dumps(dupes)}, accept_downloads=True)
    page.goto(URL); page.click('[data-tab="parc"]')
    page.set_input_files("#parc-file", str(parc)); page.wait_for_function("document.getElementById('parc-info').textContent.includes('lignes')", timeout=20000)
    years = page.eval_on_selector_all(".parc-years button:not(.parc-global)", "els=>els.map(e=>e.textContent)")
    check(years == [str(TODAY.year + i) for i in range(5)], f"portefeuille: 5 years from the current one ({years})")
    page.wait_for_selector(".parc-map .dep", timeout=10000)   # the map file is loaded on demand
    check(page.locator(".parc-map .dot").count() == 2 and page.locator(".parc-map .dep").count() == 96, "portefeuille: map of the 96 départements, one dot per signed plant")
    page.hover(".parc-map .dot[data-dep='30']"); page.wait_for_timeout(100)
    check(page.is_visible(".parc-map-wrap .parc-map-tip") and "CS DU MOULIN" in page.inner_text(".parc-map-wrap .parc-map-tip") and "Solaire PV" in page.inner_text(".parc-map-wrap .parc-map-tip"), "portefeuille: hovering a dot shows the plant and its technology")
    page.click(".parc-map .dot[data-dep='30']"); page.wait_for_timeout(100)
    check("gard" in page.inner_text(".parc-map-side").lower(), "portefeuille: a plant is placed where it is (Département2), not at the contract's département")
    check(page.inner_text(".parc-map-side .map-row.two .map-p") == "SOLARIS", "portefeuille: each installation of the département shows its producer")
    page.click(".parc-map-side .map-plant"); page.wait_for_timeout(150)
    check(page.is_visible("#co-sheet-overlay") and page.inner_text("#co-sheet-title") == "Solaris" and "CS DU MOULIN" in page.inner_text("#co-sheet-overlay"),
          "map: a plant of the open département opens its client's sheet, with the contract")
    page.click("#co-sheet-close"); page.click('[data-tab="parc"]'); page.wait_for_timeout(100)
    vb0 = page.get_attribute(".parc-map", "viewBox")
    page.click(".map-zoom button >> nth=0"); page.wait_for_timeout(100)
    check(page.get_attribute(".parc-map", "viewBox") != vb0, "portefeuille: the map zooms in")
    for _ in range(3): page.click(".map-zoom button >> nth=0")
    page.wait_for_function("document.querySelector('.parc-map').dataset.hd==='true'", timeout=10000)
    check(len(page.get_attribute(".parc-map .reg-line", "d")) > 150000 and "zoomed" in page.get_attribute(".parc-map", "class"),
          "portefeuille: zoomed in, the detailed outlines are loaded and the shadow is dropped (sharp borders)")
    page.click(".map-zoom button >> nth=2"); page.wait_for_timeout(100)
    check(page.get_attribute(".parc-map", "viewBox").startswith("-14 -14 "), "portefeuille: reset shows all France again")
    def centre(code):
        bb = page.locator(f".parc-map path.dep[data-dep='{code}']").bounding_box(); return bb["x"] + bb["width"] / 2, bb["y"] + bb["height"] / 2
    x, y = centre("29"); page.mouse.move(x, y)
    for _ in range(3): page.mouse.wheel(0, -100); page.wait_for_timeout(60)
    x, y = centre("29"); page.mouse.click(x, y); page.wait_for_timeout(200)
    vb_zoomed = page.get_attribute(".parc-map", "viewBox")
    check("finistère" in page.inner_text(".parc-map-side .map-eyebrow").lower() and not vb_zoomed.startswith("-14 -14 "),
          "portefeuille: on a zoomed map, a click on a département opens it (the press used to pan the map and swallow the click)")
    x, y = centre("22"); page.mouse.click(x, y); page.wait_for_timeout(200)
    check("finistère" in page.inner_text(".parc-map-side .map-eyebrow").lower() and page.get_attribute(".parc-map", "viewBox") == vb_zoomed,
          "portefeuille: a click on a département without plants keeps the selection and the zoom")
    page.click(".map-zoom button >> nth=2"); page.click(".map-back"); page.wait_for_timeout(100)
    page.hover(".parc-kpis .stat-card >> nth=1"); page.wait_for_timeout(100)
    check(page.is_visible(".parc-kpis .chart-tip") and "Solaire PV" in page.inner_text(".parc-kpis .chart-tip"), "portefeuille: hovering a figure shows its detail")
    page.click(".parc-kpis .stat-card >> nth=2"); page.wait_for_timeout(150)
    check(page.is_visible(".parc-kpi-filter") and page.locator(".parc-map .dot").count() == 1 and page.locator(".parc-table tbody tr").count() == 1,
          "portefeuille: clicking « fins de contrat » filters the map and the table")
    page.click(".parc-kpi-filter button"); page.wait_for_timeout(150)
    check(not page.is_visible(".parc-kpi-filter") and page.locator(".parc-map .dot").count() == 2, "portefeuille: removing the filter shows everything again")
    page.click("#parc-source-btn"); page.fill("#parc-link", "https://example.sharepoint.com/sites/x/Suivi.xlsx"); page.click("#parc-link-save")
    check(page.get_attribute("#parc-open-link", "href") == "https://example.sharepoint.com/sites/x/Suivi.xlsx" and page.is_visible("#parc-open-link"), "portefeuille: Excel link saved, opens in one click")
    page.hover(".parc-chart .hz >> nth=0"); page.wait_for_timeout(100)
    check(page.is_visible(".parc-chart-wrap .chart-tip") and "GWh" in page.inner_text(".parc-chart-wrap .chart-tip"), "portefeuille: hovering a year of the chart shows its detail")
    page.click(".parc-years .parc-global"); k = page.inner_text(".parc-kpis")
    check("50,5 GWh" in k and "%d–%d" % (TODAY.year, TODAY.year + 4) in k, f"portefeuille: global view sums the 5 years ({k[:70]!r})")
    page.click(".parc-years button:text-is('2026')")
    k = page.inner_text(".parc-kpis")
    check("27,5 GWh" in k and "18 %" in k, f"parc: volume of won contracts vs objective ({k[:60]!r})")
    check("CS DU MOULIN" in page.inner_text("#parc-body") and page.locator(".parc-due").count() == 1, "parc: contract ending soon is listed")
    renew = [t for t in stored(page, "todo-widget-tasks") if t["id"].startswith("renew_")]
    check(len(renew) == 1 and "CS DU MOULIN" in renew[0]["title"] and renew[0]["urgency"] == "haute", "parc: renewal task created for the contract ending soon")
    page.set_input_files("#parc-file", str(parc)); page.wait_for_timeout(1500)
    check(len([t for t in stored(page, "todo-widget-tasks") if t["id"].startswith("renew_")]) == 1, "parc: reloading the file doesn't duplicate the renewal task")
    page.click(".parc-due .pq-link:has-text('Remettre dans le pipeline')"); page.wait_for_timeout(300)
    check(page.is_visible("#opp-detail-overlay") and "CS DU MOULIN" in page.inner_text("#opp-detail-title"), "parc: contract put back into the pipeline")
    page.click("#opp-detail-actions >> text=Fermer")
    page.click('[data-tab="contacts"]'); page.click("#vcf-btn")
    check(page.inner_text("#vcf-new") == "1", "vcf: duplicates inside the app count once")
    with page.expect_download() as dl:
        page.click("#vcf-go")
    card = pathlib.Path(dl.value.path()).read_text()
    check(card.count("BEGIN:VCARD") == 1 and "TEL;TYPE=CELL:0611223344" in card and "anne@urba.fr" in card, "vcf: merged card keeps phone and email")
    page.click("#vcf-close"); page.click("#vcf-btn")
    check(page.inner_text("#vcf-new") == "0", "vcf: a second export only sends new contacts")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Portefeuille synced through the private data repo: loaded on the computer, read on the phone
    def repo_meta(page, private):
        page.route("**/api.github.com/repos/o/r", lambda r: r.fulfill(status=200, content_type="application/json", body=json.dumps({"private": private})))
    files = {"tasks.json": [], "crm.json": {"companies": [], "contacts": [], "meetings": [], "opportunities": []}, "notes.json": []}
    page, errors = new_page(browser, SYNC_ON, viewport={"width": 1300, "height": 850})
    mock_github(page, files); repo_meta(page, True)
    page.goto(URL); page.wait_for_timeout(500); page.click('[data-tab="parc"]')
    page.set_input_files("#parc-file", str(parc)); page.wait_for_function("document.getElementById('parc-info').textContent.includes('lignes')", timeout=20000)
    page.wait_for_timeout(1500)
    check("parc.json" in files and len(files["parc.json"]["items"]) == 3, "portefeuille sync: the workbook loaded on the computer goes to the private data repo")
    check(not errors, f"no page errors ({errors})")
    page.close()
    phone, errors = new_page(browser, SYNC_ON, viewport={"width": 390, "height": 800}, is_mobile=True, has_touch=True)
    mock_github(phone, files); repo_meta(phone, True)
    phone.goto(URL + "?lite"); phone.click("#lite-nav [data-go=parc]")
    phone.wait_for_function("document.getElementById('parc-info').textContent.includes('3 lignes')", timeout=10000)
    check(phone.locator(".parc-kpis .stat-card").count() >= 3 and "synchronisé" in phone.inner_text("#parc-info"), "portefeuille sync: the phone shows the Portefeuille without loading any file")
    check(phone.evaluate("document.documentElement.scrollWidth") <= 392, "portefeuille sync: fits the phone width")
    check(not errors, f"no page errors ({errors})")
    phone.close()
    files2 = {"tasks.json": [], "crm.json": {"companies": [], "contacts": [], "meetings": [], "opportunities": []}, "notes.json": []}
    page, errors = new_page(browser, SYNC_ON, viewport={"width": 1300, "height": 850})
    mock_github(page, files2); repo_meta(page, False)
    page.goto(URL); page.wait_for_timeout(500); page.click('[data-tab="parc"]')
    page.set_input_files("#parc-file", str(parc)); page.wait_for_function("document.getElementById('parc-info').textContent.includes('lignes')", timeout=20000)
    page.wait_for_timeout(1500); page.click('[data-tab="lists"]'); page.click('[data-tab="parc"]')
    check("parc.json" not in files2 and "gardé sur cet appareil" in page.inner_text("#parc-info"), "portefeuille sync: never sent if GitHub does not confirm the repo is private")
    page.close()

    # --- Pipeline ↔ Suivi Agrégation: each deal sent to the Suivi says whether it is there, with the right status
    sv_opps = [{"id": "sv1", "kind": "installation", "company": "Solaris", "producteur": "SOLARIS", "centrale": "CS du Moulin", "title": "CS du Moulin", "stage": "onboarding",
                "kwc": "2990", "createdAt": 1, "updatedAt": 1, "deleted": False},
               {"id": "sv2", "kind": "installation", "company": "X", "producteur": "X", "centrale": "Toit", "title": "Toit", "stage": "onboarding", "kwc": "500", "createdAt": 1, "updatedAt": 1, "deleted": False},
               {"id": "sv3", "kind": "installation", "company": "Nouveau", "centrale": "Ferme des Pins", "title": "Ferme des Pins", "stage": "negociation", "kwc": "300",
                "createdAt": 1, "updatedAt": 1, "deleted": False},
               {"id": "sv4", "kind": "installation", "company": "Nouveau", "centrale": "Hangar", "title": "Hangar", "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(sv_opps)}, viewport={"width": 1400, "height": 900})
    page.goto(URL); page.click('[data-tab="parc"]')
    page.set_input_files("#parc-file", str(parc)); page.wait_for_function("document.getElementById('parc-info').textContent.includes('lignes')", timeout=20000)
    page.click('[data-tab="pipeline"]')
    badge = lambda name: page.locator(".pipe-card", has_text=name).locator(".suivi-badge")
    check(badge("CS du Moulin").get_attribute("class").endswith("ok") and badge("Toit").get_attribute("class").endswith("diff")
          and badge("Ferme des Pins").get_attribute("class").endswith("missing") and badge("Hangar").count() == 0,
          "pipeline ↔ Suivi: in the Suivi, status to update, missing — and nothing before « Demande de prix »")
    page.click(".pipe-card:has-text('Toit')")
    check("Suivi : En attente → Gagné" in page.inner_text("#opp-detail-body"), "pipeline ↔ Suivi: the sheet says which status to put in the Suivi")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Excel bridge (local PowerShell server, mocked here): typed rows sent, status shown
    sent = []
    def bridge(route, request):
        if request.method == "OPTIONS":
            return route.fulfill(status=204, headers={"Access-Control-Allow-Origin": "*", "Access-Control-Allow-Headers": "*", "Access-Control-Allow-Methods": "*"})
        h = {"Access-Control-Allow-Origin": "*", "Content-Type": "application/json"}
        if request.url.endswith("/status"):
            ok = request.headers.get("x-bridge-token") == "k1"
            return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "version": "1.2", "authorized": ok, "pipe": {"name": "Pipe.xlsx", "found": True}, "bdd": {"name": "Suivi.xlsx", "found": True}}))
        sent.append((request.url.rsplit("/", 2)[-2], json.loads(request.post_data)))
        return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "results": [{"ok": True, "action": "ajoutée", "row": 12}]}))
    inst = [{"id": "i1", "kind": "installation", "company": "SCEA", "producteur": "SCEA", "centrale": "Les Vignes", "title": "Les Vignes", "stage": "demande_prix",
             "goDate": "2026-10-19", "dateEntree": "2026-10-12", "kwc": "2400", "new": True, "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(inst)})
    page.route("http://127.0.0.1:8765/**", bridge)
    page.goto(URL); page.click('[data-tab="lists"]'); page.fill("#bridge-token", "k1"); page.click("#bridge-save-btn")
    page.wait_for_function("document.getElementById('bridge-status').textContent.includes('Connecté')", timeout=5000)
    check(True, "bridge: status shows the connected files")
    page.click('[data-tab="pipeline"]'); page.click(".pipe-card:has-text('Les Vignes')")
    page.click("#opp-detail-actions >> text=Envoyer au Pipe Excel"); page.wait_for_timeout(300)
    page.click("#opp-detail-actions >> text=Envoyer à la BDD Excel"); page.wait_for_timeout(300)
    pipe_row = next(b["rows"][0]["values"] for k, b in sent if k == "pipe"); bdd_row = next(b["rows"][0]["values"] for k, b in sent if k == "bdd")
    check(len(pipe_row) == 37 and pipe_row[3] == {"d": "2026-10-12"} and pipe_row[22] == 2400 and pipe_row[1] is True and pipe_row[30] == "Transféré BDD",
          "bridge: Pipe row sent with real dates, numbers and booleans")
    check(len(bdd_row) == 91 and bdd_row[3] == {"d": "2026-10-19"} and bdd_row[33] == "Les Vignes" and bdd_row[24] == "En préparation", "bridge: BDD row sent")
    check("ajoutée" in page.inner_text("#bridge-msg"), "bridge: result shown")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- Excel bridge, shared file: a cell changed by a colleague is shown, nothing replaced without a choice
    sent2 = []
    def bridge2(route, request):
        if request.method == "OPTIONS":
            return route.fulfill(status=204, headers={"Access-Control-Allow-Origin": "*", "Access-Control-Allow-Headers": "*", "Access-Control-Allow-Methods": "*"})
        h = {"Access-Control-Allow-Origin": "*", "Content-Type": "application/json"}
        if request.url.endswith("/status"):
            return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "version": "1.2", "authorized": True, "pipe": {"name": "Pipe.xlsx", "found": True}, "bdd": {"name": "Suivi.xlsx", "found": True}}))
        body = json.loads(request.post_data); sent2.append(body)
        if len(sent2) == 1:
            res = {"ok": True, "action": "mise à jour", "row": 12, "key": "SCEA / Les Vignes", "cells": 2, "written": [3, 22], "mine": [0, 3, 22],
                   "conflicts": [{"i": 30, "header": "Stade", "excel": "Perdu", "app": "Transféré BDD", "comment": False},
                                 {"i": 35, "header": "Commentaires", "excel": "Vu au salon (Marie)", "app": "Relancer lundi", "comment": True}]}
        else:
            res = {"ok": True, "action": "mise à jour", "row": 12, "key": "SCEA / Les Vignes", "cells": 1, "written": [35], "mine": [35], "conflicts": []}
        return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "results": [res], "fileWasClosed": len(sent2) == 1}))
    inst2 = [dict(inst[0], commentaires="Relancer lundi")]
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(inst2), "todo_bridge-config": json.dumps({"port": 8765, "token": "k1"})})
    page.route("http://127.0.0.1:8765/**", bridge2)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click(".pipe-card:has-text('Les Vignes')")
    page.click("#opp-detail-actions >> text=Envoyer au Pipe Excel")
    page.wait_for_selector("#xl-review-overlay:not(.hidden)", timeout=5000)
    first = sent2[0]["rows"][0]
    review = page.inner_text("#xl-review-overlay")
    check(len(first["values"]) == 37 and first["prev"] is None and "Vu au salon (Marie)" in review and "Perdu" in review
          and "ouvre-le avant" in page.inner_text("#bridge-msg"), "shared Excel: changed cells shown, closed file reported")
    opts = page.eval_on_selector_all("#xl-review-body select", "els=>els.map(e=>[...e.options].map(o=>o.value).join(','))")
    check(opts == ["keep,replace", "keep,replace,merge"], f"shared Excel: « Fusionner » offered for comments only ({opts})")
    page.select_option("#xr-0-35", "merge"); page.click("#xl-review-apply")
    page.wait_for_function("document.getElementById('xl-review-overlay').classList.contains('hidden')")
    page.wait_for_timeout(300)
    second = sent2[1]["rows"][0] if len(sent2) > 1 else {}
    o = stored(page, "todo-widget-opportunities")[0]
    check(second.get("seen") == {"35": "Vu au salon (Marie)"}, "shared Excel: the value seen is sent with the choice (replaced only if still there)")
    check(second.get("force") == {"35": "merge"} and second.get("target") == 12 and o["excelKeep"]["pipe"] == {"30": "Perdu"},
          f"shared Excel: merge sent for the comment, the kept status remembered ({second.get('force')}, {o.get('excelKeep')})")
    check(o["excelPrev"]["pipe"][3] == {"d": "2026-10-12"} and o["excelPrev"]["pipe"][35] == "Relancer lundi" and o["excelPrev"]["pipe"][30] is None,
          "shared Excel: the app remembers what it wrote, cell by cell (never the kept ones)")
    check(not errors, f"no page errors ({errors})")
    page.close()
    # an older bridge (could overwrite colleagues' cells) is refused until updated
    posts = []
    def old_bridge(route, request):
        if request.method == "OPTIONS":
            return route.fulfill(status=204, headers={"Access-Control-Allow-Origin": "*", "Access-Control-Allow-Headers": "*", "Access-Control-Allow-Methods": "*"})
        h = {"Access-Control-Allow-Origin": "*", "Content-Type": "application/json"}
        if request.url.endswith("/status"):
            return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "version": "1.1", "authorized": True, "pipe": {"name": "Pipe.xlsx", "found": True}, "bdd": {"name": "Suivi.xlsx", "found": True}}))
        posts.append(request.url); return route.fulfill(status=200, headers=h, body=json.dumps({"ok": True, "results": []}))
    page, errors = new_page(browser, {"todo_todo-widget-opportunities": json.dumps(inst), "todo_bridge-config": json.dumps({"port": 8765, "token": "k1"})})
    page.route("http://127.0.0.1:8765/**", old_bridge)
    page.goto(URL); page.click('[data-tab="pipeline"]'); page.click(".pipe-card:has-text('Les Vignes')")
    page.click("#opp-detail-actions >> text=Envoyer au Pipe Excel"); page.wait_for_timeout(600)
    check(not posts and "ancienne version" in page.inner_text("#bridge-msg"), "an older bridge is refused: nothing sent, update asked")
    page.close()

    # --- Réf. & data: clients grouped by company type, collapsed by default, search opens them
    cos = [{"id": f"co{i}", "name": n, "type": t, "whatTheyDo": "", "marketScope": [], "website": "", "createdAt": 1, "updatedAt": 1, "deleted": False}
           for i, (n, t) in enumerate([("Alpha", "Fournisseur"), ("Beta", "Client"), ("Gamma", "Client"), ("Delta", "")])]
    opps_b = [{"id": "ob1", "kind": "installation", "company": "Beta", "producteur": "Beta", "title": "Toit", "stage": "collecte", "createdAt": 1, "updatedAt": 1, "deleted": False}]
    page, errors = new_page(browser, {"todo_todo-widget-clients": json.dumps(["Alpha", "Beta", "Gamma", "Delta"]),
                                      "todo_todo-widget-companies": json.dumps(cos), "todo_todo-widget-opportunities": json.dumps(opps_b)})
    page.goto(URL); page.click('[data-tab="lists"]')
    heads = page.eval_on_selector_all("#client-list-rows .client-type-head", "els=>els.map(e=>e.textContent)")
    check(heads == ["Client· 2", "Fournisseur· 1", "Sans type· 1"], f"client sections by type {heads}")
    check(page.locator("#client-list-rows .client-list-row:visible").count() == 0, "sections collapsed by default")
    page.locator("#client-list-rows .client-type-head", has_text="Client").click()
    check(page.locator("#client-list-rows .client-list-row:visible").count() == 2, "opening a section shows its clients")
    page.reload(); page.click('[data-tab="lists"]')
    check(page.locator("#client-list-rows .client-list-row:visible").count() == 2, "open sections remembered")
    page.fill("#client-list-search", "alp")
    check(page.locator("#client-list-rows .client-list-row:visible").count() == 1, "search opens matching sections")
    page.fill("#client-list-search", "")
    page.once("dialog", lambda d: d.accept("Beta Énergie"))
    page.locator("#client-list-rows .client-list-row", has_text="Beta").locator("button[title=Renommer]").click()
    opp = next(o for o in stored(page, "todo-widget-opportunities") if o["id"] == "ob1")
    check(opp["company"] == "Beta Énergie" and opp["producteur"] == "Beta Énergie", f"renaming a client renames its opportunities too ({opp['company']})")
    check(not errors, f"no page errors ({errors})")
    page.close()

    # --- XSS: conflict values are shown as text
    conflicts = [{"id": "cf1", "entityType": "task", "entityId": "t1", "entityName": "T", "field": "title", "fieldLabel": "Titre",
                  "localValue": "<img src=x onerror=\"window.__xss=1\">", "remoteValue": "ok", "resolvedValue": "ok"}]
    page, errors = new_page(browser, {"todo_sync-conflicts": json.dumps(conflicts)})
    page.goto(URL); page.click('[data-tab="lists"]')
    check(page.evaluate("window.__xss") is None, "conflict panel doesn't execute HTML")
    page.close()

    # --- Service worker (needs http): GitHub reads are never served from the cache
    import functools, http.server, threading
    root = pathlib.Path(__file__).resolve().parent.parent
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(root)))
    srv.RequestHandlerClass.log_message = lambda *a: None
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    ctx = browser.new_context(service_workers="allow")
    reads = {"n": 0}
    def gh(route, req):
        reads["n"] += 1
        route.fulfill(status=200, content_type="application/json", headers={"Access-Control-Allow-Origin": "*"},
                      body=json.dumps({"sha": "s%d" % reads["n"], "encoding": "base64", "content": base64.b64encode(b"[]").decode()}))
    ctx.route("**/api.github.com/**", gh)
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    page = ctx.new_page(); page.goto("http://127.0.0.1:%d/index.html" % srv.server_address[1])
    page.wait_for_function("navigator.serviceWorker && navigator.serviceWorker.controller", timeout=15000)
    shas = [page.evaluate("fetch('https://api.github.com/repos/o/r/contents/data/tasks.json').then(r=>r.json()).then(d=>d.sha)") for _ in range(3)]
    check(shas == ["s1", "s2", "s3"], f"service worker never serves sync data from its cache ({shas})")
    ctx.close(); srv.shutdown()

    # --- Service worker: a new version of the page shows at the first reload (pages are network first)
    import shutil
    site = pathlib.Path(tempfile.mkdtemp()) / "site"
    shutil.copytree(root, site, ignore=shutil.ignore_patterns(".git", "tests", "__pycache__"))
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(site)))
    srv.RequestHandlerClass.log_message = lambda *a: None
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    ctx = browser.new_context(service_workers="allow")
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    page = ctx.new_page(); page.goto("http://127.0.0.1:%d/index.html" % srv.server_address[1])
    page.wait_for_function("navigator.serviceWorker && navigator.serviceWorker.controller", timeout=15000)
    html = (site / "index.html").read_text(encoding="utf-8")
    (site / "index.html").write_text(html.replace("</body>", "<div id='new-version-marker'></div></body>"), encoding="utf-8")
    page.reload()
    check(page.locator("#new-version-marker").count() == 1, "a published update shows at the first reload")
    app_js = (site / "js" / "app.js").read_text(encoding="utf-8")
    (site / "js" / "app.js").write_text(app_js + "\ndocument.body.dataset.codeVersion='new';\n", encoding="utf-8")
    page.reload(); page.wait_for_timeout(300)
    check(page.evaluate("document.body.dataset.codeVersion") == "new", "a published update of the app's code (js/app.js) runs at the first reload too")
    # a new service worker version takes over by itself and the open page reloads once onto it
    sw = (site / "sw.js").read_text(encoding="utf-8")
    (site / "sw.js").write_text(sw.replace("todo-pwa-v", "todo-pwa-test-v"), encoding="utf-8")
    (site / "index.html").write_text(html.replace("</body>", "<div id='second-version-marker'></div></body>"), encoding="utf-8")
    page.evaluate("navigator.serviceWorker.getRegistration().then(r=>r.update())")
    page.wait_for_selector("#second-version-marker", state="attached", timeout=15000)
    check(page.locator("#second-version-marker").count() == 1, "an open page switches to a newly published version by itself")
    ctx.close(); srv.shutdown()

    browser.close()

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
